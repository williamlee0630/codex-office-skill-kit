"""離線 CSV / XLSX 描述統計；不推論因果、不執行來源公式。"""
import argparse
import csv
import math
import statistics
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from _office import require, safe_cell, write_json, write_markdown, write_xlsx, protect_input


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def number(value):
    if blank(value) or isinstance(value, (bool, date, datetime)):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return result if math.isfinite(result) else None


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if not isinstance(value, str):
        return None
    try:
        return date.fromisoformat(value.strip()).isoformat()
    except ValueError:
        return None


def load_table(path, sheet=None):
    require(path.is_file(), '找不到輸入資料。')
    if path.suffix.lower() == '.csv':
        try:
            with path.open(encoding='utf-8-sig', newline='') as stream:
                table = list(csv.reader(stream))
        except UnicodeError as exc:
            raise ValueError('CSV 請另存為 UTF-8 編碼。') from exc
        source_sheet = None
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        require(not sheet or sheet in wb.sheetnames, '找不到指定工作表。')
        ws = wb[sheet] if sheet else wb.active
        source_sheet = ws.title
        table = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValueError('僅支援 CSV 與 XLSX。舊版 XLS 請先另存 XLSX。')
    require(bool(table), '資料為空，請提供欄位名稱與至少一筆資料。')
    headers = [str(v).strip() if v is not None else '' for v in table[0]]
    require(all(headers) and len(headers) == len(set(headers)), '欄位名稱不可空白或重複。')
    rows = []
    row_numbers = []
    for row_no, row in enumerate(table[1:], 2):
        if all(blank(v) for v in row):
            continue
        require(len(row) <= len(headers), f'第 {row_no} 列的欄位數超過標題列，請檢查 CSV 分隔符號。')
        rows.append(row + [None] * (len(headers) - len(row)))
        row_numbers.append(row_no)
    require(bool(rows), '資料只有標題或空白列，無法計算統計。')
    return headers, rows, row_numbers, source_sheet


def summarize(values, row_numbers):
    nonempty = [(i, v) for i, v in zip(row_numbers, values) if not blank(v)]
    nums = [(i, number(v)) for i, v in nonempty if number(v) is not None]
    dates = [(i, iso_date(v)) for i, v in nonempty if iso_date(v) is not None]
    result = {'type': 'text', 'missing_count': len(values) - len(nonempty), 'nonempty_count': len(nonempty)}
    if dates and len(dates) / max(1, len(nonempty)) >= 0.6:
        result.update(type='date', invalid_count=len(nonempty) - len(dates),
                      min_date=min(v for _, v in dates), max_date=max(v for _, v in dates))
    elif nums and len(nums) / max(1, len(nonempty)) >= 0.6:
        series = [n for _, n in nums]
        outliers = []
        # 四分位距採 inclusive 插值；少於四個有效值不標記異常。
        if len(series) >= 4:
            q1, _, q3 = statistics.quantiles(series, n=4, method='inclusive')
            iqr = q3 - q1
            if iqr > 0:
                outliers = [i for i, n in nums if n < q1 - 1.5 * iqr or n > q3 + 1.5 * iqr]
        result.update(type='number', count=len(series), invalid_count=len(nonempty) - len(nums),
                      invalid_rows=[i for i, v in nonempty if number(v) is None],
                      sum=math.fsum(series), mean=statistics.mean(series), median=statistics.median(series),
                      min=min(series), max=max(series),
                      sample_stdev=statistics.stdev(series) if len(series) > 1 else None, outlier_rows=outliers)
    else:
        result.update(distinct_count=len({str(v) for _, v in nonempty}), invalid_count=0)
    return result


def grouped(rows, key_index, value_index, dates=False):
    buckets = defaultdict(list)
    omitted = []
    for row_no, row in rows:
        key = iso_date(row[key_index]) if dates else str(row[key_index]).strip() if not blank(row[key_index]) else None
        value = number(row[value_index])
        if key is None or value is None:
            omitted.append(row_no)
            continue
        buckets[key].append(value)
    return {k: {'count': len(v), 'sum': math.fsum(v), 'mean': statistics.mean(v)} for k, v in sorted(buckets.items())}, omitted


def draw_chart(path, labels, values, title, line=False):
    from PIL import Image, ImageDraw, ImageFont
    require(len(labels) == len(values) and bool(values), '圖表沒有可用數值。')
    image = Image.new('RGB', (1200, 720), '#F5F8F8')
    draw = ImageDraw.Draw(image)
    font_path = Path('C:/Windows/Fonts/msjh.ttc')
    font = ImageFont.truetype(str(font_path), 21) if font_path.exists() else ImageFont.load_default(size=21)
    heading = ImageFont.truetype(str(font_path), 30) if font_path.exists() else ImageFont.load_default(size=30)
    draw.text((55, 28), title, fill='#123C3B', font=heading)
    left, top, right, bottom = 100, 130, 1140, 560
    lower, upper = min(0, min(values)), max(0, max(values))
    span = upper - lower or 1
    project = lambda v: bottom - (v - lower) / span * (bottom - top)
    baseline = project(0)
    draw.line((left, top, left, bottom), fill='#657B7A', width=2)
    draw.line((left, baseline, right, baseline), fill='#657B7A', width=2)
    for fraction in (0, .25, .5, .75, 1):
        value = lower + fraction * span
        y = project(value)
        draw.line((left, y, right, y), fill='#DDE5E4', width=1)
        draw.text((8, y - 12), f'{value:.3g}', fill='#385251', font=font)
    width = (right - left) / len(values)
    points = []
    for i, (label, value) in enumerate(zip(labels, values)):
        x, y = left + width * (i + .5), project(value)
        points.append((x, y))
        if not line:
            draw.rectangle((x - width * .3, min(y, baseline), x + width * .3, max(y, baseline) + 1), fill='#187D78')
        # 圖表上限 12 組；完整資料仍保存在 JSON 與 Excel。
        draw.text((x - min(70, len(str(label)) * 7), bottom + 20), str(label)[:10], fill='#385251', font=font)
        draw.text((x - 25, max(top, y - 30)), f'{value:.4g}', fill='#123C3B', font=font)
    if line and len(points) > 1:
        draw.line(points, fill='#187D78', width=5)
    if line:
        for x, y in points:
            draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill='#187D78')
    draw.text((55, 655), '來源：輸入資料；僅描述觀察值，無法推論因果', fill='#385251', font=font)
    image.save(path)


def excel_report(path, headers, rows, result):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    summary_rows = []
    for column, stats in result['columns'].items():
        summary_rows.append([column, stats['type'], stats['missing_count'], stats.get('invalid_count'), stats.get('count'), stats.get('sum'), stats.get('mean'), stats.get('median'), stats.get('sample_stdev'), ', '.join(map(str, stats.get('outlier_rows', [])))])
    write_xlsx(path, ['欄位', '型態', '缺漏數', '非有效數', '有效數值數', '合計', '平均', '中位數', '樣本標準差', 'IQR 異常來源列'], summary_rows, sheet_name='分析摘要')
    wb = load_workbook(path)
    source = wb.create_sheet('原始資料')
    source.append([safe_cell(v) for v in headers])
    for row in rows:
        source.append([safe_cell(v) for v in row])
    source.append([])
    source.append(['來源檔案', safe_cell(result['source'])])
    source.append(['計算說明', '合計為 Python 計算的靜態結果；修改原始資料後請重新執行。'])
    source.append(['資料日期欄', safe_cell(result.get('date_column') or '未指定，待確認')])
    source.append(['限制', '異常標記需人工確認；描述統計無法證明因果。'])
    for name, mapping in [('分組統計', result['groups']), ('日期趨勢', {d['date']: d for d in result['trend']})]:
        if mapping:
            ws = wb.create_sheet(name)
            ws.append(['分組' if name == '分組統計' else '日期', '有效筆數', '合計', '平均'])
            for key, entry in mapping.items():
                ws.append([safe_cell(key), entry['count'], entry['sum'], entry['mean']])
            chart = LineChart() if name == '日期趨勢' else BarChart()
            chart.title = f"{result['value_column']} 合計"
            chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
            chart.width, chart.height = 22, 12
            ws.add_chart(chart, 'F2')
    for ws in wb:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='164E4A')
            cell.font = Font(name='Microsoft JhengHei', bold=True, color='FFFFFF')
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 25 if ws.title != '原始資料' else 24
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Microsoft JhengHei', size=11)
                if isinstance(cell.value, (int, float)) and ws.title != '原始資料':
                    cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
    wb.save(path)


def analyze(args):
    protect_input(args.input, args.output)
    path, out = Path(args.input), Path(args.output)
    require(not any((out / name).exists() for name in ('analysis.xlsx', 'analysis.json', 'analysis.md', 'chart.png')),
            '輸出資料夾已有成果，請使用新的輸出資料夾，避免舊圖表或報告被誤認為本次結果。')
    headers, rows, row_numbers, sheet = load_table(path, args.sheet)
    columns = {header: summarize([r[i] for r in rows], row_numbers) for i, header in enumerate(headers)}
    for requested in (args.group_by, args.value, args.date_column):
        require(not requested or requested in headers, f'找不到指定欄位：{requested}')
    require(not (args.group_by or args.date_column) or args.value, '分組或日期趨勢需要 --value 指定數值欄位。')
    if args.value:
        require(columns[args.value]['type'] == 'number', '指定值欄位未達數值比例 60%，請先核對型態。')
    result = {'source': str(path.resolve()), 'source_sheet': sheet, 'row_count': len(rows), 'columns': columns,
              'duplicate_row_count': len(rows) - len({tuple(map(str, r)) for r in rows}),
              'groups': {}, 'trend': [], 'group_column': args.group_by, 'value_column': args.value,
              'date_column': args.date_column, 'group_omitted_rows': [], 'trend_omitted_rows': [],
              'method': 'Python statistics；數值欄須至少 60% 非空值可轉為有限數值；IQR inclusive，至少四筆且 IQR > 0；空白列排除。',
              'limitations': ['靜態計算，修改資料後須重新執行。', '缺漏不補零；異常列仍保留。', '描述趨勢不代表因果關係。']}
    if args.group_by:
        result['groups'], result['group_omitted_rows'] = grouped(zip(row_numbers, rows), headers.index(args.group_by), headers.index(args.value))
    if args.date_column:
        trend, result['trend_omitted_rows'] = grouped(zip(row_numbers, rows), headers.index(args.date_column), headers.index(args.value), dates=True)
        result['trend'] = [dict(date=k, **v) for k, v in trend.items()]
    numeric = [name for name, stats in columns.items() if stats['type'] == 'number']
    out.mkdir(parents=True, exist_ok=True)
    require(path.resolve() not in {(out / 'analysis.xlsx').resolve(), (out / 'analysis.json').resolve()}, '輸出不可覆寫輸入檔案。')
    chart_info = '沒有數值欄位，未產生圖表。'
    if result['trend']:
        entries = result['trend'][:12]
        draw_chart(out / 'chart.png', [d['date'] for d in entries], [d['sum'] for d in entries], f"{args.value} 日期合計（前 12 日期）", line=True)
        chart_info = '日期合計折線圖。單一日期只有觀察點，無法判定趨勢。' if len(entries) == 1 else '日期合計折線圖，顯示前 12 個日期，完整數值見 Excel。'
    elif result['groups']:
        entries = list(result['groups'].items())[:12]
        draw_chart(out / 'chart.png', [k for k, _ in entries], [v['sum'] for _, v in entries], f"{args.value} 分組合計（前 12 組）")
        chart_info = '分組合計長條圖，顯示前 12 組，完整數值見 Excel。'
    elif numeric:
        name = numeric[0]
        series = [number(r[headers.index(name)]) for r in rows]
        series = [v for v in series if v is not None]
        if min(series) == max(series):
            labels, counts = [f'{series[0]:g}'], [len(series)]
        else:
            lo, hi = min(series), max(series)
            width = (hi - lo) / min(8, len(series))
            counts = [0] * min(8, len(series))
            for v in series:
                counts[min(len(counts) - 1, int((v - lo) / width))] += 1
            labels = [f'{lo + i * width:.3g}' for i in range(len(counts))]
        draw_chart(out / 'chart.png', labels, counts, f'{name} 分布（橫軸為區間下界，縱軸為筆數）')
        chart_info = f'{name} 的等寬區間次數分布；最後區間包含最大值。'
    result['chart_note'] = chart_info
    write_json(out / 'analysis.json', result)
    excel_report(out / 'analysis.xlsx', headers, rows, result)
    lines = ['# 資料分析報告', '', f'來源：{path.name}；工作表：{sheet or "CSV"}。資料日期：{args.date_column or "未指定，待確認"}。',
             f'有效資料列：{len(rows)}。完全相同的重複列：{result["duplicate_row_count"]}（保留且納入計算）。', '',
             '## 描述統計', '', '| 欄位 | 型態 | 缺漏 | 非有效值 | 合計 | 平均 | 中位數 |', '| --- | --- | ---: | ---: | ---: | ---: | ---: |']
    for name, stats in columns.items():
        lines.append(f'| {name.replace("|", "／")} | {stats["type"]} | {stats["missing_count"]} | {stats.get("invalid_count", 0)} | {stats.get("sum", "不適用")} | {stats.get("mean", "不適用")} | {stats.get("median", "不適用")} |')
    for name, stats in columns.items():
        if stats.get('outlier_rows'):
            lines.append(f'\n{name} 的 IQR 異常來源列：{stats["outlier_rows"]}，需人工確認，未刪除。\n')
    lines.extend(['', '## 圖表與比較', '', chart_info, '', '分組及日期合計只納入有效數值與有效分組／日期；排除列保存在 analysis.json。',
                  '', '## 方法與待確認', '', result['method'], '數字識別不自動判斷欄位的業務意義。識別碼即使全為數字也不應當作績效。',
                  '本報告只支持描述性觀察，不能推論因果。Codex 應依原始資料核對異常原因、幣別、資料期間及分組意義。',
                  'Excel 為靜態結果，修改資料後請重新執行。公式來源保留為安全文字，未運算其內容。'])
    write_markdown(out / 'analysis.md', '\n'.join(lines))
    print('已產生分析、Excel 與可用圖表：' + str(out.resolve()))


def main():
    parser = argparse.ArgumentParser(description='離線計算 CSV / XLSX 描述統計，不執行來源公式。')
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--sheet', help='XLSX 工作表名稱；省略使用啟用的工作表。')
    parser.add_argument('--group-by', help='分組欄位名稱。')
    parser.add_argument('--value', help='分組與趨勢的數值欄位名稱。')
    parser.add_argument('--date-column', help='ISO 日期（YYYY-MM-DD）欄位名稱。')
    args = parser.parse_args()
    try:
        analyze(args)
        return 0
    except ImportError:
        print('錯誤：缺少 openpyxl 或 Pillow，請安裝工具包 requirements.txt。', file=sys.stderr)
    except Exception as exc:
        print('錯誤：無法完成資料分析。' + str(exc), file=sys.stderr)
    return 1


if __name__ == '__main__':
    sys.exit(main())
