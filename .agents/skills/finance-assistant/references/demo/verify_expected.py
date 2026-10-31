"""教學模擬資料的獨立驗算；不匯入或執行財務正式腳本。"""
from collections import Counter, defaultdict
import csv
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
D = Decimal


def expected(folder):
    return json.loads((folder / 'expected.json').read_text(encoding='utf-8'))


def rounded(value):
    return format(value.quantize(D('.01'), rounding=ROUND_HALF_UP), '.2f')


def verify_reconciliation():
    folder = ROOT / 'reconciliation' if (ROOT / 'reconciliation').is_dir() else ROOT
    exp = expected(folder)
    def read(name):
        with (folder / name).open(encoding='utf-8-sig', newline='') as f:
            return {n: r for n, r in enumerate(csv.DictReader(f), 2)}
    internal, bank = read('internal.csv'), read('bank.csv')
    counts = Counter()
    used_internal, used_bank = set(), set()
    for left, right, status in exp['pairs']:
        a, b = internal[left], bank[right]
        assert left not in used_internal and right not in used_bank
        used_internal.add(left); used_bank.add(right)
        assert a['currency'] == b['currency'] and a['date'][:7] == b['date'][:7]
        if status == '已確認匹配':
            assert a['reference'] == b['reference'] != '' and D(a['amount']) == D(b['amount'])
            assert a['date'] == b['date']
        elif status == '金額不符':
            assert a['reference'] == b['reference'] != ''
            assert rounded(D(a['amount']) - D(b['amount'])) == exp['mismatch_internal_minus_bank']
        else:
            assert D(a['amount']) == D(b['amount']) and a['counterparty'] == b['counterparty']
            assert (a['reference'] == b['reference'] == '' and a['date'] == b['date']) or abs((date.fromisoformat(a['date']) - date.fromisoformat(b['date'])).days) > 3
        counts[status] += 1
    for row in exp['duplicate_internal_rows']:
        assert internal[row]['reference'] == 'DUP' and D(internal[row]['amount']) == D('50')
        assert row not in used_internal
        used_internal.add(row); counts['疑似重複紀錄'] += 1
    for row in exp['other_missing_rows']:
        assert row not in used_internal
        used_internal.add(row); counts['銀行或另一來源缺漏'] += 1
    for row in exp['internal_missing_rows']:
        assert row not in used_bank
        used_bank.add(row); counts['內部帳務缺漏'] += 1
    assert used_internal == set(internal) and used_bank == set(bank)
    assert dict(counts) == exp['counts']
    # 單獨相同金額而無對象、跨期資料不屬於已確認交易。
    assert internal[11]['amount'] == bank[9]['amount'] and not internal[11]['counterparty']
    assert internal[12]['reference'] == bank[10]['reference'] and internal[12]['date'][:7] != bank[10]['date'][:7]


def verify_budget():
    folder = ROOT / 'budget'; exp = expected(folder)
    wb = load_workbook(folder / 'budget-actual.xlsx', data_only=False)
    rows = list(wb.active.values)[1:]; wb.close()
    differences, percentages = [], []
    totals = defaultdict(lambda: [D('0'), D('0'), D('0')])
    for item, base, actual, currency, period, note in rows:
        assert '教學模擬' in note and period == '2026-10'
        if base is None or actual is None:
            differences.append(None); percentages.append(None); continue
        base, actual = D(str(base)), D(str(actual))
        difference = actual - base
        differences.append(rounded(difference))
        percentages.append(rounded(difference * 100 / base) if base else None)
        totals[currency][0] += base; totals[currency][1] += actual; totals[currency][2] += difference
    assert differences == exp['differences']
    assert percentages == exp['percentages']
    assert {currency: dict(zip(['base', 'actual', 'difference'], map(rounded, values)))
            for currency, values in totals.items()} == exp['totals']


def verify_trial_balance():
    folder = ROOT / 'trial-balance'; exp = expected(folder)['group']
    wb = load_workbook(folder / 'trial-balance.xlsx', data_only=False)
    rows = list(wb.active.values)[1:]; wb.close()
    debit = sum((D(r[2]) for r in rows), D(0))
    credit = sum((D(r[3]) for r in rows), D(0))
    net = defaultdict(lambda: D(0))
    for account, classification, debit_row, credit_row, currency, period, note in rows:
        assert currency == 'TWD' and period == '2026-10' and '教學模擬' in note
        # 獨立以借方淨額一次累計，再從分類恆等式驗算。
        net[classification] += D(debit_row) - D(credit_row)
    profit = -net['收入'] - net['費用']
    assert debit == credit == D('450000')
    assert rounded(debit) == exp['debit_total'] and rounded(credit) == exp['credit_total']
    assert rounded(profit) == exp['profit']
    assert rounded(net['資產']) == exp['assets']
    assert rounded(-net['負債']) == exp['liabilities']
    assert rounded(-net['權益'] + profit) == exp['equity']
    assert rounded(net['資產'] + net['負債'] + net['權益'] - profit) == exp['balance_difference']


if __name__ == '__main__':
    verify_reconciliation(); verify_budget(); verify_trial_balance()
    print('三組教學模擬資料的獨立來源驗算通過；未呼叫正式財務計算腳本。')
