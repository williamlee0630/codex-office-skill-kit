"""獨立技能共用執行資源；由 sync_runtime.py 複製，無網路與發布功能。"""
from pathlib import Path
import argparse
import json
import math
import sys
import unicodedata


def require(condition, message):
    if not condition:
        raise ValueError(message)


def protect_input(input_path, output_dir):
    """阻止來源與已知產物同名，保留原始資料。"""
    names = {'evidence.json', 'layout-check.json', 'analysis.json', 'receipts-summary.json',
             'extraction-evidence.json', 'transcription-template.json', 'minutes.md', 'minutes.docx',
             'progress.md', 'progress.csv', 'progress.xlsx', 'weekly-brief.md', 'weekly-detailed.md',
             'weekly-brief.docx', 'weekly-detailed.docx', 'document-formal.md', 'document-concise.md',
             'document-formal.docx', 'document-concise.docx', 'ideas.md', 'comparison.csv',
             'selected-plan.md', 'daily-summary.md', 'drafts.md', 'research-report.md', 'sources.csv',
             'analysis.xlsx', 'analysis.md', 'chart.png', 'presentation.pptx', 'speaker-notes.md',
             'receipts.xlsx', 'receipts-report.md', 'review-required.md'}
    source = Path(input_path).resolve()
    out = Path(output_dir).resolve()
    require(all(source != out / name for name in names), '輸入檔案與產物路徑相同，請改用新的輸出資料夾，以保留原始資料。')


def load_json(path):
    path = Path(path)
    require(path.suffix.lower() == '.json', '需要 JSON 檔案；請先依來源整理成技能指定的結構。')
    try:
        data = json.loads(path.read_text(encoding='utf-8-sig'),
                          parse_constant=lambda value: (_ for _ in ()).throw(ValueError('數值必須為有限值：' + value)))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError(f'無法讀取 JSON：{path.name}。請確認 UTF-8 編碼與 JSON 格式。') from exc
    require(isinstance(data, dict), 'JSON 最外層必須是物件，不能是陣列或空值。')
    return data


def read_text(path):
    path = Path(path)
    require(path.is_file(), f'找不到來源檔案：{path}')
    ext = path.suffix.lower()
    try:
        if ext in {'.txt', '.md', '.csv', '.json'}:
            text = path.read_text(encoding='utf-8-sig')
        elif ext == '.docx':
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc = Document(path)
            blocks = []
            for element in doc.element.body.iterchildren():
                if element.tag.endswith('}p'):
                    blocks.append(Paragraph(element, doc).text)
                elif element.tag.endswith('}tbl'):
                    blocks.extend(' | '.join(cell.text for cell in row.cells) for row in Table(element, doc).rows)
            text = '\n'.join(blocks)
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            require(not reader.is_encrypted, 'PDF 已加密；請提供可讀取且已授權的副本。')
            text = '\n'.join(f'[第 {i + 1} 頁]\n{page.extract_text() or ""}'
                             for i, page in enumerate(reader.pages))
            # 頁碼標籤不算成功辨識的文字。
            require(any((page.extract_text() or '').strip() for page in reader.pages),
                    'PDF 沒有可擷取文字。請使用 Codex 可用的視覺/OCR 工具或人工轉錄，保留待確認欄位。')
        else:
            raise ValueError(f'不支援 {ext or "無副檔名"}。可讀 TXT、MD、CSV、JSON、DOCX、文字型 PDF；圖片需視覺辨識或人工轉錄。')
    except ImportError as exc:
        raise ValueError('缺少格式套件，請依 docs/INSTALL.md 安裝 requirements.txt。') from exc
    except (OSError, UnicodeError) as exc:
        raise ValueError('無法讀取檔案；文字檔請另存 UTF-8，並確認檔案沒有被鎖定。') from exc
    require(bool(text.strip()), '來源檔案沒有文字，請提供內容後再執行。')
    return text


def write_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    return path


def write_markdown(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(str(text).rstrip() + '\n', encoding='utf-8')
    return path


def safe_cell(value):
    if isinstance(value, float):
        require(math.isfinite(value), '試算表不能寫入 NaN 或無限值。')
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def write_xlsx(path, headers, rows, sheet_name='明細'):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append([safe_cell(h) for h in headers])
    for row in rows:
        ws.append([safe_cell(v) for v in row])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(name='Microsoft JhengHei', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='164E4A')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Microsoft JhengHei', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(52, max(18, len(str(header)) * 2 + 6))
    for row in ws.iter_rows(min_row=2):
        needed = 1
        for cell in row:
            width = ws.column_dimensions[cell.column_letter].width - 2
            line_count = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1 for c in line) / width))
                             for line in str(cell.value or '').split('\n'))
            needed = max(needed, line_count)
        ws.row_dimensions[row[0].row].height = min(409, max(24, needed * 17 + 10))
    wb.save(path)
    return path


def write_docx(path, title, sections):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml.ns import qn
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = 'AI 辦公室工具包'
    doc.core_properties.title = str(title)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.1)
    sec.left_margin = sec.right_margin = Cm(2.3)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Microsoft JhengHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.add_paragraph(str(title), style='Title')
    for heading, body in sections:
        if heading:
            doc.add_heading(str(heading), level=1)
        paragraphs = body if isinstance(body, (list, tuple)) else str(body).split('\n')
        for paragraph in paragraphs:
            doc.add_paragraph(str(paragraph))
    doc.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description='擷取本機文件文字，不執行 OCR 或網路存取。')
    parser.add_argument('command', choices=['extract'])
    parser.add_argument('input')
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        source, target = Path(args.input), Path(args.output)
        same_file = source.resolve() == target.resolve() or (source.exists() and target.exists() and source.samefile(target))
        require(not same_file, '擷取輸出不可指向原始檔案，請另選新的文字檔以保留來源。')
        write_markdown(args.output, read_text(args.input))
        print('文字已擷取：' + str(Path(args.output).resolve()))
        return 0
    except (ValueError, OSError, ImportError) as exc:
        print('無法完成：' + str(exc), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
