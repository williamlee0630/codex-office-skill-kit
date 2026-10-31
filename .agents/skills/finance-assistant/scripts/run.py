"""本機財務計算：直接讀取原始表格，以 Decimal 運算；不修改或過帳來源。"""
import argparse
from collections import Counter, defaultdict
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import math
import unicodedata
from zipfile import BadZipFile

from _office import require, safe_cell, write_docx, write_json

ZERO = Decimal('0')
CURRENCIES = set('TWD USD EUR JPY CNY HKD GBP AUD CAD CHF SGD NZD KRW THB MYR IDR VND PHP INR SEK NOK DKK ZAR BRL MXN AED SAR BHD KWD CLP'.split())
DIGITS = {'JPY': 0, 'KRW': 0, 'VND': 0, 'CLP': 0, 'BHD': 3, 'KWD': 3}
ALIASES = {
    'date': ['日期', '交易日期'], 'reference': ['交易編號', '憑證編號', '參考號碼', '交易識別碼'],
    'amount': ['金額', '交易金額'], 'currency': ['幣別', '貨幣'], 'period': ['期間', '會計期間'],
    'counterparty': ['交易對象', '對象', '收付款對象'], 'account': ['科目', '會計科目', '科目名稱'],
    'class': ['分類', '科目分類'], 'debit': ['借方', '借方金額'], 'credit': ['貸方', '貸方金額'],
    'balance': ['餘額'], 'direction': ['借貸方向'], 'item': ['項目', '項目名稱'],
    'budget': ['預算', '預算金額', 'base', '基準'], 'actual': ['實際', '實際金額', 'current', '本期'],
    'department': ['部門'], 'cash_flow_class': ['現金流量分類'], 'cash_flow_amount': ['現金流量金額'],
}
CLASSES = {'asset': 'asset', '資產': 'asset', 'liability': 'liability', '負債': 'liability',
           'equity': 'equity', '權益': 'equity', '股東權益': 'equity', 'revenue': 'revenue',
           '收入': 'revenue', '營收': 'revenue', 'expense': 'expense', '費用': 'expense', '成本': 'expense'}
STATUSES = ['已確認匹配', '疑似匹配，待人工確認', '金額不符', '內部帳務缺漏', '銀行或另一來源缺漏', '疑似重複紀錄']


def decimal(value, location):
    if value is None or str(value).strip() == '':
        return None
    text = str(value).strip().replace(',', '')
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise ValueError(f'{location} 金額不是可計算數字：{value}') from None
    require(result.is_finite(), f'{location} 金額不可為 NaN 或無限值。')
    require(len(result.as_tuple().digits) <= 28 and abs(result.adjusted()) <= 20,
            f'{location} 金額超過可驗證的 28 位精度或範圍。')
    return result


def money(value, currency='TWD'):
    if value is None:
        return None
    unit = Decimal(1).scaleb(-DIGITS.get(currency, 2))
    return format(value.quantize(unit, rounding=ROUND_HALF_UP), 'f')


def num(value):
    if value is None:
        return None
    # Excel 只保證 15 位有效數字；大額仍提供確切十進位文字，不能靜默截斷分位。
    exact = Decimal(str(value))
    significant = len(exact.normalize().as_tuple().digits)
    return exact if significant > 15 else float(value)


def read_table(path, mapping=None, sheet=None):
    path = Path(path)
    require(path.is_file(), f'找不到原始資料：{path}')
    suffix = path.suffix.lower()
    metadata = []
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False, read_only=True)
        require(sheet is None or sheet in wb.sheetnames, f'找不到工作表：{sheet}')
        ws = wb[sheet] if sheet else wb.active
        require(ws is not None, 'Excel 沒有可讀取的工作表。')
        records = []
        for row in ws.iter_rows():
            require(not any(c.data_type == 'f' for c in row),
                    f'{path.name} 含公式；請另存經確認的數值快照，不能使用過期快取。')
            records.append((row[0].row, [c.value for c in row]))
        sheet_name = ws.title
        wb.close()
    elif suffix in {'.csv', '.txt', '.tsv', '.md'}:
        text = path.read_text(encoding='utf-8-sig')
        physical = text.splitlines()
        meaningful = [(n, line) for n, line in enumerate(physical, 1) if line.strip()]
        require(meaningful, f'{path.name} 沒有資料。')
        if suffix in {'.txt', '.md'}:
            metadata = [{'line': n, 'text': line} for n, line in meaningful if line.lstrip().startswith('#') and ',' not in line and '\t' not in line]
            meta_lines = {m['line'] for m in metadata}
            meaningful = [(n, line) for n, line in meaningful if n not in meta_lines]
        else: meta_lines = set()
        require(meaningful, f'{path.name} 沒有表格。')
        if meaningful[0][1].strip().startswith('|'):
            records = [(n, [x.strip() for x in line.strip().strip('|').split('|')]) for n, line in meaningful
                       if not re.fullmatch(r'[\s|:\-]+', line)]
        else:
            delimiter = '\t' if '\t' in meaningful[0][1] else ','
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            records = []
            while True:
                start_line = reader.line_num + 1
                try: record = next(reader)
                except StopIteration: break
                if start_line in meta_lines: continue
                if any(str(v).strip() for v in record): records.append((start_line, record))
        sheet_name = ''
    else:
        raise ValueError('財務原始資料支援 XLSX、CSV、TXT、TSV、Markdown 表格。')
    require(len(records) > 1, f'{path.name} 沒有明細。')
    headers = [str(v).strip() if v is not None else '' for v in records[0][1]]
    require(all(headers) and len(set(headers)) == len(headers), f'{path.name} 欄位名稱空白或重複。')
    mapping = mapping or {}
    require(isinstance(mapping, dict), 'columns 必須是標準欄位到來源欄名的對照。')
    for key, field in mapping.items():
        require(key in ALIASES and field in headers, f'欄位對照不存在：{key} → {field}')
    found = {}
    for key, aliases in ALIASES.items():
        candidates = [h for h in headers if h == key or h in aliases]
        if key in mapping: found[key] = mapping[key]
        elif len(candidates) == 1: found[key] = candidates[0]
        elif len(candidates) > 1: raise ValueError(f'{path.name} 的 {key} 有多個候選欄位，請明確指定 columns。')
    rows = []
    for index, values_row in records[1:]:
        if not any(v is not None and str(v).strip() for v in values_row): continue
        require(len(values_row) <= len(headers), f'{path.name} 第 {index} 列欄位數超過標題。')
        values_row = list(values_row) + [None] * (len(headers) - len(values_row))
        raw = dict(zip(headers, values_row))
        row = {key: raw[field] for key, field in found.items()}
        row.update(_row=index, _source=path.name, _sheet=sheet_name, _raw=raw)
        rows.append(row)
    require(rows, f'{path.name} 沒有非空白明細。')
    return {'path': str(path.resolve()), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'headers': headers, 'mapping': found, 'rows': rows, 'sheet': sheet_name, 'metadata': metadata}


def clean(value):
    return '' if value is None else str(value).strip()


def location(row):
    return f"{row['_source']}／{row['_sheet'] or '文字表格'}／第 {row['_row']} 列"


def normalize(rows, config, review, transaction=False):
    for row in rows:
        row['currency'] = clean(row.get('currency') or config.get('currency')).upper()
        row['period'] = clean(row.get('period'))
        raw_date = row.get('date')
        parsed = None
        if raw_date:
            try:
                parsed = raw_date.date() if isinstance(raw_date, datetime) else raw_date if isinstance(raw_date, date) else date.fromisoformat(clean(raw_date).replace('/', '-'))
            except ValueError:
                review.append(location(row) + '：日期無法辨識，待人工確認。')
        row['_date'] = parsed
        if not row['period']:
            row['period'] = parsed.strftime('%Y-%m') if parsed else clean(config.get('period'))
        issues = []
        if row['currency'] not in CURRENCIES: issues.append('幣別空白或未支援，需確認')
        if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])|\d{4}', row['period']): issues.append('會計期間空白或格式錯誤')
        if parsed and row['period'] and not parsed.isoformat().startswith(row['period']): issues.append('交易日期與會計期間不一致')
        if config.get('period') and row['period'] != config['period']: issues.append('與指定會計期間不一致')
        if transaction:
            row['reference'] = clean(row.get('reference'))
            row['counterparty'] = clean(row.get('counterparty'))
            row['_amount'] = decimal(row.get('amount'), location(row))
            if row['_amount'] is None and ('debit' in row or 'credit' in row):
                debit, credit = decimal(row.get('debit'), location(row)), decimal(row.get('credit'), location(row))
                if debit is not None and credit is not None: row['_amount'] = debit - credit
            if row['_amount'] is None: issues.append('金額空白（不視為零）')
            if not parsed: issues.append('缺少有效交易日期')
        row['_valid'] = not issues
        for issue in issues: review.append(location(row) + '：' + issue)


def sources_rows(tables):
    rows = []
    for table in tables:
        for row in table['rows']:
            for key, value in row['_raw'].items():
                rows.append([row['_source'], row['_row'], key, clean(value), table['sha256'], row['_sheet']])
        for item in table.get('metadata', []):
            rows.append([Path(table['path']).name, item['line'], '文字標記（非交易）', item['text'], table['sha256'], ''])
    return rows


def basis_rows(tables):
    result = [
        ['運算', 'Decimal 原始十進位金額累加；ROUND_HALF_UP 僅於輸出依幣別小數位四捨五入。'],
        ['幣別', '分幣別、期間計算；不換匯、不混算；未知幣別與缺漏列保留於原始資料及待確認。'],
        ['空白', '空白金額與零不同，不自動補零；來源公式拒讀，需經確認數值快照。'],
        ['Excel 精度', '超過 15 位有效數字的金額以確切十進位文字存放，其餘金額為數值；JSON 保留原始十進位計算值。'],
        ['限制', '輔助工作成果，未完成正式會計查核；原因無證據時列待查，不自動過帳。'],
    ]
    for table in tables:
        result.append(['來源', table['path'] + '；SHA256=' + table['sha256']])
        for field, header in table['mapping'].items(): result.append(['欄位對照', f'{table["path"]}：{field} ← {header}'])
    return result


def statements(table, config):
    review = []
    rows = table['rows']
    normalize(rows, config, review)
    grouped = defaultdict(list)
    seen = Counter()
    valid_count = 0
    for row in rows:
        cls = CLASSES.get(clean(row.get('class')))
        account = clean(row.get('account'))
        debit, credit = decimal(row.get('debit'), location(row)), decimal(row.get('credit'), location(row))
        if debit is None and credit is None and 'balance' in row:
            balance = decimal(row.get('balance'), location(row))
            direction = clean(row.get('direction')).lower()
            if balance is not None and direction in {'debit', '借', '借方', 'credit', '貸', '貸方'}:
                debit, credit = (balance, ZERO) if direction in {'debit', '借', '借方'} else (ZERO, balance)
            else: review.append(location(row) + '：餘額需要明確借貸方向，未自行推測。')
        if not account: review.append(location(row) + '：缺少會計科目。')
        if not cls: review.append(location(row) + '：科目分類缺少或不支援。')
        if debit is None or credit is None: review.append(location(row) + '：借貸金額空白，需填明確零或提供餘額與方向。')
        if not row['_valid'] or not account or not cls or debit is None or credit is None: continue
        if debit != 0 and credit != 0: review.append(location(row) + '：同列借貸均非零，可能是總帳期間累計，請核對來源。')
        key = (row['currency'], row['period'])
        signature = (key, account, cls, debit, credit)
        seen[signature] += 1
        if seen[signature] > 1: review.append(location(row) + '：疑似重複科目列；保留原始累計，不自動刪除。')
        grouped[key].append((row, cls, debit, credit))
        valid_count += 1
    income, balance_sheet, checks, groups, flows = [], [], [], [], []
    for (currency, period), items in sorted(grouped.items()):
        totals = {kind: ZERO for kind in ['asset', 'liability', 'equity', 'revenue', 'expense']}
        debits = sum((i[2] for i in items), ZERO)
        credits = sum((i[3] for i in items), ZERO)
        for _, cls, debit, credit in items:
            totals[cls] += debit - credit if cls in {'asset', 'expense'} else credit - debit
        profit = totals['revenue'] - totals['expense']
        difference = totals['asset'] - totals['liability'] - totals['equity'] - profit
        balanced = debits == credits
        if not balanced: review.append(f'{currency}／{period} 借貸不平衡，原始精度差額 {debits - credits}；未新增調整分錄。')
        if difference != ZERO: review.append(f'{currency}／{period} 資產負債表不平衡，原始精度差額 {difference}。')
        for label, value in [('收入', totals['revenue']), ('費用', totals['expense']), ('本期損益', profit)]:
            income.append([currency, period, label, num(money(value, currency))])
        for label, value in [('資產', totals['asset']), ('負債', totals['liability']), ('原始權益', totals['equity']),
                             ('本期損益', profit), ('權益含本期損益', totals['equity'] + profit), ('平衡差額', difference)]:
            balance_sheet.append([currency, period, label, num(money(value, currency))])
        checks.append([currency, period, num(money(debits, currency)), num(money(credits, currency)),
                       num(money(debits - credits, currency)), balanced, num(money(difference, currency))])
        group = {'currency': currency, 'period': period, 'debit_total': money(debits, currency),
                 'credit_total': money(credits, currency), 'debit_credit_balanced': balanced,
                 'profit': money(profit, currency), 'assets': money(totals['asset'], currency),
                 'liabilities': money(totals['liability'], currency), 'equity': money(totals['equity'] + profit, currency),
                 'balance_difference': money(difference, currency), 'cash_flow_available': False,
                 'balance_difference_exact': str(difference), 'balance_sheet_balanced': difference == ZERO,
                 'exact_values': {k: str(v) for k, v in {**totals, 'debits': debits, 'credits': credits}.items()}}
        # 試算表餘額本身不能推導現金流量；只接受逐筆且完整分類的現金金額及期初期末餘額。
        cash = config.get('cash_flow')
        if cash and currency == cash.get('currency') and period == cash.get('period'):
            movement = defaultdict(lambda: ZERO)
            complete = len(items) == len(rows)
            for row, _, _, _ in items:
                amount = decimal(row.get('cash_flow_amount'), location(row))
                kind = clean(row.get('cash_flow_class'))
                if amount is None or kind not in {'operating', 'investing', 'financing', 'noncash'}:
                    complete = False
                elif kind == 'noncash' and amount != ZERO: complete = False
                else: movement[kind] += amount
            opening = decimal(cash.get('opening_cash'), '期初現金')
            closing = decimal(cash.get('closing_cash'), '期末現金')
            complete = complete and opening is not None and closing is not None
            if complete and opening + sum(movement.values(), ZERO) == closing:
                group['cash_flow_available'] = True
                for label, value in [('期初現金', opening), ('營業活動', movement['operating']),
                                     ('投資活動', movement['investing']), ('籌資活動', movement['financing']), ('期末現金', closing)]:
                    flows.append([currency, period, label, num(money(value, currency))])
            else: review.append(f'{currency}／{period} 現金流量分類、金額或期初期末核對不足，未編製現金流量表。')
        if not group['cash_flow_available']: review.append(f'{currency}／{period} 缺少足夠現金流量交易分類及核對資料，未編製現金流量表。')
        groups.append(group)
    complete = valid_count == len(rows) and bool(groups) and all(g['debit_credit_balanced'] and g['balance_sheet_balanced'] for g in groups) and not any('重複' in r for r in review)
    review.append('會計準則及報表用途：' + clean(config.get('accounting_standard') or '未指定；屬工作底稿，正式報表適用準則與揭露需人工確認。'))
    sheets = [('損益表', ['幣別', '會計期間', '項目', '金額'], income),
              ('資產負債表', ['幣別', '會計期間', '項目', '金額'], balance_sheet),
              ('核對摘要', ['幣別', '期間', '借方合計', '貸方合計', '借貸差額', '借貸平衡', '資產負債差額'], checks)]
    if flows: sheets.append(('現金流量表', ['幣別', '期間', '項目', '金額'], flows))
    basis = [['損益', '收入＝貸方−借方；費用＝借方−貸方；本期損益＝收入−費用。'],
             ['資產負債', '資產＝借−貸；負債及原始權益＝貸−借；權益含本期損益；差額＝資產−負債−權益−本期損益。'],
             ['完整性', 'complete 僅表示輸入列可分類且平衡；不代表正式財報、揭露、現金流量或會計查核完整。']]
    return {'groups': groups, 'complete': complete, 'review': review}, sheets, basis


def reconcile(internal, other, config):
    review = []
    left, right = internal['rows'], other['rows']
    normalize(left, config, review, True)
    normalize(right, config, review, True)
    results, used_left, used_right = [], set(), set()

    def add(status, a=None, b=None, reason=''):
        currency = (a or b)['currency']
        amount_a, amount_b = a.get('_amount') if a else None, b.get('_amount') if b else None
        difference = amount_a - amount_b if amount_a is not None and amount_b is not None else None
        results.append({'status': status, 'internal_row': a['_row'] if a else None,
                        'other_row': b['_row'] if b else None, 'currency': currency,
                        'period': (a or b)['period'], 'internal_amount': money(amount_a, currency),
                        'other_amount': money(amount_b, currency), 'difference': money(difference, currency),
                        'reason': reason, 'reference': (a or b).get('reference', '')})
        if a: used_left.add(a['_row'])
        if b: used_right.add(b['_row'])
        if status != STATUSES[0]: review.append(f'{status}：內部第 {a["_row"] if a else "—"} 列／另一來源第 {b["_row"] if b else "—"} 列；{reason}')

    # 重複交易識別碼（同來源、幣別、期間）全部隔離；無識別碼用完整交易特徵，絕不用相同金額單獨去重。
    for rows, side in [(left, 'internal'), (right, 'other')]:
        signatures = defaultdict(list)
        for row in rows:
            if row.get('reference'):
                key = ('ref', row['currency'], row['period'], row['reference'])
            else:
                key = ('row', row['currency'], row['period'], row.get('_date'), row.get('_amount'), row.get('counterparty'), json.dumps(row['_raw'], sort_keys=True, default=str))
            signatures[key].append(row)
        for same in signatures.values():
            if len(same) > 1:
                for row in same:
                    add(STATUSES[5], row if side == 'internal' else None, row if side == 'other' else None,
                        '同來源交易識別碼或完整列重複；全部保留並排除自動配對。')

    def compatible(a, b):
        return a['_valid'] and b['_valid'] and a['currency'] == b['currency'] and a['period'] == b['period']

    tolerance = config.get('date_tolerance_days', 3)
    require(isinstance(tolerance, int) and 0 <= tolerance <= 31, '日期容許天數需為 0 到 31 的整數。')
    for a in left:
        if a['_row'] in used_left or not a['reference']: continue
        candidates = [b for b in right if b['_row'] not in used_right and compatible(a, b) and a['reference'] == b['reference']]
        if len(candidates) == 1:
            b = candidates[0]
            if a['_amount'] != b['_amount']:
                add(STATUSES[2], a, b, '同一交易識別碼、幣別與期間，金額不同；差額＝內部−另一來源。')
            elif abs((a['_date'] - b['_date']).days) <= tolerance and (not a['counterparty'] or not b['counterparty'] or a['counterparty'] == b['counterparty']):
                add(STATUSES[0], a, b, '唯一交易識別碼、金額、幣別、期間相符，日期在容許範圍；一對一配對。')
            else: add(STATUSES[1], a, b, '識別碼金額相符，但日期差異或交易對象衝突，需人工確認。')
    # 無相同識別碼，只將雙向唯一的日期＋交易對象＋金額候選列為疑似，不能因金額相同確認。
    def weak(a, b):
        return compatible(a, b) and a['_amount'] == b['_amount'] and a['_date'] == b['_date'] and bool(a['counterparty']) and a['counterparty'] == b['counterparty']
    for a in left:
        if a['_row'] in used_left: continue
        candidates = [b for b in right if b['_row'] not in used_right and weak(a, b)]
        if len(candidates) == 1:
            b = candidates[0]
            reverse = [x for x in left if x['_row'] not in used_left and weak(x, b)]
            if len(reverse) == 1: add(STATUSES[1], a, b, '日期、對象與金額相符但缺少一致交易識別碼；僅疑似。')
    for a in left:
        if a['_row'] not in used_left: add(STATUSES[4], a, reason='另一來源未找到可安全一對一配對交易；缺漏資料須先確認。')
    for b in right:
        if b['_row'] not in used_right: add(STATUSES[3], b=b, reason='內部未找到可安全一對一配對交易；缺漏資料須先確認。')
    counts = {status: sum(r['status'] == status for r in results) for status in STATUSES}
    # 報表按交易來源列排序，便於回查；已配對列不再重複使用。
    results.sort(key=lambda r: (r['internal_row'] is None, r['internal_row'] or r['other_row']))
    sheet_rows = [[r['status'], r['internal_row'], r['other_row'], r['currency'], num(r['internal_amount']),
                   num(r['other_amount']), num(r['difference']), r['reason'], r['period'], r['reference']] for r in results]
    totals = []
    for table, side in [(internal, '內部'), (other, '另一來源')]:
        grouped = defaultdict(lambda: ZERO)
        for row in table['rows']:
            if row['_valid']: grouped[(row['currency'], row['period'])] += row['_amount']
        for (currency, period), value in sorted(grouped.items()): totals.append([side, currency, period, num(money(value, currency)), '包含原始重複列；不自動刪除'])
    return {'results': results, 'counts': counts, 'review': review, 'complete': counts[STATUSES[0]] == len(results)}, [
        ('分析結果', ['核對結果', '內部原始列', '另一來源原始列', '幣別', '內部金額', '另一來源金額', '差額', '判斷依據', '期間', '交易識別碼'], sheet_rows),
        ('分幣期間合計', ['來源', '幣別', '期間', '合計', '限制'], totals)], [
        ['對帳差額', '內部金額−銀行或另一來源金額；沖銷負數保留正負號。'],
        ['確定配對條件', f'同幣別同期間，唯一識別碼、相同金額、日期差不超過 {tolerance} 天；對象不得衝突。'],
        ['一對一', '每一來源列至多使用一次；重複識別碼全部隔離；日期對象金額候選也只能列疑似。']]


def variance(table, other, config):
    review = []
    normalize(table['rows'], config, review)
    if other: normalize(other['rows'], config, review)
    results = []
    comparison = config.get('comparison', 'budget')
    require(comparison in {'budget', 'month', 'year', 'department'}, 'comparison 支援 budget、month、year、department。')

    def add(row, base, actual, base_period, actual_period, other_row=None, note=''):
        currency = row['currency']
        item = clean(row.get('item') or row.get('account'))
        if not item:
            review.append(location(row) + '：缺少項目名稱。')
            note += '缺少項目；'
        valid = row['_valid'] and bool(item) and not note
        if base is None or actual is None:
            review.append(location(row) + '：基準或實際金額空白，未視為零。')
            valid = False
        difference = actual - base if valid else None
        percentage = difference / base * Decimal(100) if valid and base != ZERO else None
        if valid and base == ZERO: review.append(location(row) + '：比較基準為零，百分比無法直接計算。')
        abnormal = valid and (base == ZERO and actual != ZERO or percentage is not None and abs(percentage) >= Decimal(str(config.get('anomaly_percent', 20))))
        direction = '增加' if difference is not None and difference > 0 else '減少' if difference is not None and difference < 0 else '持平' if difference == 0 else '待確認'
        results.append({'item': item, 'currency': currency, 'base_period': base_period, 'actual_period': actual_period,
                        'base': money(base, currency), 'actual': money(actual, currency), 'difference': money(difference, currency),
                        'percentage': money(percentage, 'TWD'), 'direction': direction, 'abnormal': bool(abnormal),
                        'source_row': row['_row'], 'other_row': other_row, 'reason': note or ('變動原因待查；未提供原因證據。' if abnormal else ''),
                        'exact_base': str(base) if base is not None else None, 'exact_actual': str(actual) if actual is not None else None,
                        'exact_difference': str(difference) if difference is not None else None})

    if other:
        # 兩來源比較：source 永遠是比較基準；item＋幣別＋部門為配對鍵，跨月份/年度時保留雙方期間。
        def key(row): return (clean(row.get('item') or row.get('account')), row['currency'], '' if comparison == 'department' else clean(row.get('department')))
        left = defaultdict(list); right = defaultdict(list)
        for row in table['rows']: left[key(row)].append(row)
        for row in other['rows']: right[key(row)].append(row)
        for item_key in sorted(set(left) | set(right)):
            a, b = left[item_key], right[item_key]
            if len(a) == 1 and len(b) == 1:
                base_row, actual_row = a[0], b[0]
                note = ''
                if not actual_row['_valid']: note = '另一來源欄位缺漏，待確認。'
                if comparison == 'budget' and base_row['period'] != actual_row['period']: note = '預算與實際期間不同，未混算。'
                if comparison == 'department' and base_row['period'] != actual_row['period']: note = '部門比較期間不同，未混算。'
                add(base_row, decimal(base_row.get('amount'), location(base_row)), decimal(actual_row.get('amount'), location(actual_row)),
                    base_row['period'], actual_row['period'], actual_row['_row'], note)
            else:
                for row in a + b:
                    review.append(location(row) + '：比較項目缺少另一來源或有多個候選列，未自行合併。')
                    add(row, decimal(row.get('amount'), location(row)) if row in a else None,
                        decimal(row.get('amount'), location(row)) if row in b else None,
                        row['period'] if row in a else '', row['period'] if row in b else '', note='比較來源不唯一或缺漏。')
    else:
        seen = Counter((clean(r.get('item') or r.get('account')), r['currency'], r['period'], clean(r.get('department'))) for r in table['rows'])
        for row in table['rows']:
            signature = (clean(row.get('item') or row.get('account')), row['currency'], row['period'], clean(row.get('department')))
            note = ''
            if seen[signature] > 1:
                note = '疑似重複比較項目；未自動合併。'
                review.append(location(row) + '：' + note)
            add(row, decimal(row.get('budget'), location(row)), decimal(row.get('actual'), location(row)), row['period'], row['period'], note=note)
    totals = defaultdict(lambda: {'base': ZERO, 'actual': ZERO, 'difference': ZERO, 'rows': 0})
    for row in results:
        if row['difference'] is not None:
            key = (row['currency'], row['base_period'], row['actual_period'])
            totals[key]['base'] += Decimal(row['exact_base']); totals[key]['actual'] += Decimal(row['exact_actual'])
            totals[key]['difference'] += Decimal(row['exact_difference']); totals[key]['rows'] += 1
    totals_result = [{'currency': c, 'base_period': p, 'actual_period': q,
                      **{k: money(v, c) if k != 'rows' else v for k, v in value.items()}}
                     for (c, p, q), value in sorted(totals.items())]
    sheet_rows = [[r['item'], r['currency'], num(r['base']), num(r['actual']), num(r['difference']),
                   num(r['percentage']), r['direction'], r['abnormal'], r['base_period'], r['actual_period'], r['source_row'], r['other_row'], r['reason']] for r in results]
    total_rows = [[r['currency'], r['base_period'], r['actual_period'], num(r['base']), num(r['actual']), num(r['difference']), r['rows']] for r in totals_result]
    return {'results': results, 'totals': totals_result, 'review': review, 'complete': all(r['difference'] is not None for r in results)}, [
        ('分析結果', ['項目', '幣別', '預算或比較基準', '實際或本期', '差額', '差異百分比（%）', '增減方向', '異常', '基準期間', '實際期間', '來源列', '另一來源列', '原因或限制'], sheet_rows),
        ('分幣期間合計', ['幣別', '基準期間', '實際期間', '基準合計', '實際合計', '差額合計', '有效列數'], total_rows)], [
        ['差額', '實際或本期−預算或比較基準；兩來源時 source 為分母基準、other 為比較值。'],
        ['差異百分比', '(實際−比較基準)÷比較基準×100；零基準保留空白，負基準保留原始分母符號。'],
        ['異常', f'百分比絕對值達 {config.get("anomaly_percent", 20)}%，或基準零而實際非零；原因需資料支持。'],
        ['彙總', '按幣別及雙方期間累加未四捨五入的有效原始金額；缺漏或重複候選列不混入。']]


def write_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title); ws.append(headers)
        for row in rows: ws.append([str(v) if isinstance(v, Decimal) else safe_cell(v) for v in row])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name='Microsoft JhengHei', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='164E4A')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 22 if len(header) < 9 else 38
        if title == '原始資料':
            ws.column_dimensions['D'].width = 42
            ws.column_dimensions['E'].width = 68
        if title in {'計算依據', '待人工確認'}:
            ws.column_dimensions['B'].width = 90
        if title == '分析結果':
            for column in ['H', 'M']: ws.column_dimensions[column].width = 48
        for row in ws.iter_rows(min_row=2):
            currency = ''
            amount_columns, integer_columns, percent_columns = set(), set(), set()
            if title in {'損益表', '資產負債表', '現金流量表'}:
                currency = row[0].value; amount_columns = {4}
            elif title == '核對摘要':
                currency = row[0].value; amount_columns = {3, 4, 5, 7}
            elif title == '分析結果' and headers[0] == '核對結果':
                currency = row[3].value; amount_columns = {5, 6, 7}; integer_columns = {2, 3}
            elif title == '分析結果':
                currency = row[1].value; amount_columns = {3, 4, 5}; integer_columns = {11, 12}; percent_columns = {6}
            elif title == '分幣期間合計' and headers[0] == '來源':
                currency = row[1].value; amount_columns = {4}
            elif title == '分幣期間合計':
                currency = row[0].value; amount_columns = {4, 5, 6}; integer_columns = {7}
            elif title == '原始資料': integer_columns = {2}
            elif title == '待人工確認': integer_columns = {1}
            digits = DIGITS.get(currency, 2)
            pattern = '#,##0' + ('.' + '0' * digits if digits else '')
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.font = Font(name='Microsoft JhengHei', size=11)
                if cell.column in amount_columns: cell.number_format = pattern + ';[Red]-' + pattern
                elif cell.column in integer_columns: cell.number_format = '0'
                elif cell.column in percent_columns: cell.number_format = '0.00"%"'
            lines = 2
            for cell in row:
                width = max(8, ws.column_dimensions[cell.column_letter].width - 2)
                for line in str(cell.value if cell.value is not None else '').split('\n'):
                    length = sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1 for c in line)
                    lines = max(lines, math.ceil(length / width))
            ws.row_dimensions[row[0].row].height = min(409, max(32, lines * 17 + 8))
    wb.save(path)


def execute(config, output, base_dir):
    mode = config.get('mode')
    require(mode in {'statements', 'reconcile', 'variance'}, '請指定財務任務：statements、reconcile 或 variance。')
    require(config.get('source'), '請提供原始帳務表格 source。')
    paths = [Path(config['source'])]
    if config.get('other'): paths.append(Path(config['other']))
    paths = [(base_dir / p).resolve() if not p.is_absolute() else p.resolve() for p in paths]
    filename = {'statements': '財務報表.xlsx', 'reconcile': '銀行對帳結果.xlsx', 'variance': '財務差異分析.xlsx'}[mode]
    output = Path(output).resolve()
    targets = [output / name for name in [filename, '財務分析報告.docx', 'finance-result.json']]
    require(all(not p.exists() for p in targets), '輸出檔已存在，請使用新的輸出資料夾；不覆寫既有成果。')
    require(all(p not in paths for p in targets), '輸出不可覆寫原始資料。')
    tables = [read_table(paths[0], config.get('columns'), config.get('sheet'))]
    if len(paths) > 1: tables.append(read_table(paths[1], config.get('other_columns'), config.get('other_sheet')))
    if mode == 'statements': result, sheets, basis = statements(tables[0], config)
    elif mode == 'reconcile':
        require(len(tables) == 2, '對帳需要內部帳務及銀行或另一來源兩份資料。')
        result, sheets, basis = reconcile(tables[0], tables[1], config)
    else: result, sheets, basis = variance(tables[0], tables[1] if len(tables) > 1 else None, config)
    result.update(mode=mode, simulation=bool(config.get('simulation')), source_row_count=sum(len(t['rows']) for t in tables),
                  sources=[{k: v for k, v in table.items() if k != 'rows'} for table in tables],
                  formal_audit_completed=False)
    if result['simulation']: result['review'].insert(0, '教學模擬資料，非真實公司帳務。')
    sheets.extend([('原始資料', ['來源檔案', '原始列', '來源欄位', '原始值', '來源SHA256', '工作表'], sources_rows(tables)),
                   ('計算依據', ['項目', '依據與限制'], basis_rows(tables) + basis),
                   ('待人工確認', ['編號', '待確認事項'], [[n, item] for n, item in enumerate(result['review'], 1)] or [[1, '未發現輸入層級問題；正式會計查核尚未完成。']])])
    # 所有驗證與計算先完成，才建立輸出；失敗不留下看似成功的財報。
    output.mkdir(parents=True, exist_ok=True)
    write_workbook(targets[0], sheets)
    summary = []
    if mode == 'statements':
        summary = [f"{g['currency']}／{g['period']}：收入減費用＝{g['profit']}；借貸合計 {g['debit_total']}／{g['credit_total']}；資產負債差額 {g['balance_difference']}。" for g in result['groups']]
    elif mode == 'reconcile': summary = [f'{k}：{v} 筆' for k, v in result['counts'].items()]
    else: summary = [f"{r['item']}／{r['currency']}：基準 {r['base'] if r['base'] is not None else '待確認'}、實際 {r['actual'] if r['actual'] is not None else '待確認'}、差額 {r['difference'] if r['difference'] is not None else '待確認'}、百分比 {r['percentage'] if r['percentage'] is not None else '無法直接計算或資料缺漏'}。" for r in result['results']]
    write_docx(targets[1], '財務分析報告' + ('（教學模擬資料）' if result['simulation'] else ''), [
        ('分析範圍與來源', [f'{t["path"]}；SHA256：{t["sha256"]}' for t in tables]),
        ('計算摘要', summary or ['資料不足，僅保留來源與待確認項目。']),
        ('計算依據', [f'{a}：{b}' for a, b in basis]),
        ('完整性與限制', ['輸入列核對狀態：' + ('可計算且通過本次輸入層級檢查' if result['complete'] else '尚有缺漏或差異，僅為部分工作底稿'),
                       '正式會計查核尚未完成；未新增交易、調整帳本或補造原因；所有明細見同名工作簿。']),
        ('待人工確認', result['review'] or ['正式報表用途與覆核程序仍需確認。'])])
    write_json(targets[2], result)
    return targets


def main():
    parser = argparse.ArgumentParser(description='本機財務報表、對帳及差異分析；保留原始資料。')
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--mode', choices=['statements', 'reconcile', 'variance'])
    parser.add_argument('--other')
    parser.add_argument('--period')
    parser.add_argument('--currency')
    args = parser.parse_args()
    try:
        source = Path(args.input).resolve()
        if source.suffix.lower() == '.json':
            config = json.loads(source.read_text(encoding='utf-8-sig'))
            require(isinstance(config, dict), '內部任務設定必須是物件。')
        else: config = {'source': str(source)}
        for name in ['mode', 'other', 'period', 'currency']:
            value = getattr(args, name)
            if value is not None: config[name] = value
        with localcontext() as context:
            context.prec = 60
            outputs = execute(config, args.output, source.parent)
        for output in outputs: print('已產生：' + str(output))
        return 0
    except (ValueError, OSError, ImportError, InvalidOperation, KeyError, TypeError, BadZipFile, csv.Error) as exc:
        print('錯誤：' + str(exc), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
