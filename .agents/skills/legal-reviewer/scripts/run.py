"""本機來源定位、引用核對與法務審查產檔。語意／法律判斷由 Codex 完成。"""
import argparse
from datetime import date
import hashlib
from pathlib import Path
import sys
import tempfile
from urllib.parse import urlparse

from _office import load_json, read_text, require, write_docx, write_json, write_markdown, write_xlsx, safe_cell

TOPICS = {
    'parties_purpose': '當事人與目的', 'term_renewal': '期限與自動續約',
    'payment': '付款條件與期限', 'liability': '違約、賠償與責任上限',
    'confidentiality_privacy': '保密與個人資料', 'ip': '智慧財產權',
    'termination': '終止條件', 'law_disputes': '適用法律與爭議解決', 'special': '特殊條款與其他權益'}
PRODUCTS = ['合約審查報告.docx', '合約風險清單.xlsx', '條款修訂建議.md', 'audit.json']
LIMIT = '本報告為 AI 初步審閱，風險評級屬輔助判斷，不等同違法認定，也不是律師正式法律意見。原文存在核對不代表語意或法律結論已獲證實。'


def text(value, label):
    require(isinstance(value, str) and bool(value.strip()), label + '需要非空文字。')
    return value


def strings(value, label):
    require(isinstance(value, list) and all(isinstance(v, str) and v.strip() for v in value), label + '需要文字陣列。')
    return value


def fresh_output(output, names, sources=()):
    out = Path(output).resolve()
    for name in names:
        target = out / name
        require(not target.exists(), f'產物已存在：{name}。請改用新的輸出資料夾，避免覆寫舊成果或原件。')
        for source in sources:
            require(Path(source).resolve() != target.resolve(), '產物與原始文件路徑相同，請改用新的輸出資料夾。')
    return out


def extract_source(source, base):
    require(isinstance(source, dict) and (bool(source.get('path')) != bool(source.get('text'))), 'source 請提供 path 或 text，兩者擇一。')
    warnings = []
    blocks = []
    path = None
    if source.get('path'):
        path = (Path(base) / text(source['path'], '來源路徑')).resolve()
        require(path.is_file(), '找不到原始文件：' + str(path))
        raw = path.read_bytes()
        ext = path.suffix.lower()
        if ext in {'.txt', '.md'}:
            content = read_text(path)
            blocks = [{'location': f'L{i}', 'text': line} for i, line in enumerate(content.splitlines(), 1) if line.strip()]
        elif ext == '.docx':
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc = Document(path)
            paragraph = table = 0
            for element in doc.element.body.iterchildren():
                if element.tag.endswith('}p'):
                    paragraph += 1
                    value = Paragraph(element, doc).text
                    if value.strip():
                        blocks.append({'location': f'P{paragraph}', 'text': value})
                elif element.tag.endswith('}tbl'):
                    table += 1
                    for r, row in enumerate(Table(element, doc).rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            if cell.text.strip():
                                blocks.append({'location': f'T{table}/R{r}/C{c}', 'text': cell.text})
            for section_num, section in enumerate(doc.sections, 1):
                for label, part in [('HEADER', section.header), ('FOOTER', section.footer)]:
                    for i, p in enumerate(part.paragraphs, 1):
                        if p.text.strip():
                            blocks.append({'location': f'S{section_num}/{label}/P{i}', 'text': p.text})
            warnings.append('DOCX 擷取正文、表格及頁首頁尾文字；文字方塊、圖片、註解、修訂追蹤與附件內容須另外檢視，未當作已審查。')
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            require(not reader.is_encrypted, 'PDF 已加密；請提供已授權且可讀取的副本。')
            for page_num, page in enumerate(reader.pages, 1):
                content = page.extract_text() or ''
                require(bool(content.strip()), f'第 {page_num} 頁沒有可擷取文字，可能為掃描或空白頁；本工具無 OCR，請先用可用 OCR／視覺工具確認並保留頁碼後再審查，不宣稱已完整讀取。')
                blocks.extend({'location': f'PAGE{page_num}/L{i}', 'text': line} for i, line in enumerate(content.splitlines(), 1) if line.strip())
            warnings.append('PDF 位置為頁碼與擷取行號，可能與畫面換行不同；圖片、閱讀順序及跨欄內容須對照原頁確認。')
        else:
            raise ValueError('僅支援 DOCX、文字型 PDF、TXT、MD 或直接貼上的文字。')
    else:
        content = text(source.get('text'), '貼上原文')
        raw = content.encode('utf-8')
        blocks = [{'location': f'L{i}', 'text': line} for i, line in enumerate(content.splitlines(), 1) if line.strip()]
    require(bool(blocks), '原始文件沒有可審查文字。')
    return {'path': str(path) if path else None, 'kind': path.suffix.lower()[1:] if path else 'pasted',
            'sha256': hashlib.sha256(raw).hexdigest(), 'blocks': blocks, 'warnings': warnings}


def citations(items, blocks, required=False):
    require(isinstance(items, list), 'citations 必須為陣列。')
    require(not required or bool(items), '現有條款必須提供可回查原文的 citations。')
    for item in items:
        require(isinstance(item, dict), '每個原文引用必須包含 location 與 quote。')
        loc, quote = item.get('location'), item.get('quote')
        require(isinstance(loc, str) and loc in blocks, '原文引用位置不存在。請重新擷取來源索引。')
        require(isinstance(quote, str) and quote.strip() and quote in blocks[loc], '原文引用與指定位置不符，請逐字核對 quote。')
    return items


def check_laws(data, base, jurisdiction):
    sources = data.get('legal_sources', [])
    require(isinstance(sources, list), 'legal_sources 必須為陣列。')
    catalog = {}
    for source in sources:
        require(isinstance(source, dict), '法律來源必須為物件。')
        sid = text(source.get('id'), '法律來源 ID')
        require(sid not in catalog, '法律來源 ID 不可重複。')
        for field in ['law_name', 'article', 'url', 'checked_at', 'jurisdiction', 'authority', 'quote']:
            text(source.get(field), '法律來源 ' + field)
        require(urlparse(source['url']).scheme in {'https', 'http'} and bool(urlparse(source['url']).netloc), '法律來源需要有效官方網址。')
        try:
            checked = date.fromisoformat(source['checked_at'])
        except ValueError as exc:
            raise ValueError('法律來源查核日期需為 YYYY-MM-DD。') from exc
        require(checked <= date.today(), '法律來源查核日期不可在未來。')
        snapshot = extract_source({'path': source.get('path')}, base)
        require(any(source['quote'] in b['text'] for b in snapshot['blocks']), '法律來源原文摘錄不符保存的來源。')
        require(source['jurisdiction'] == jurisdiction, '法律來源法域與本次審查法域不同，請分開分析。')
        catalog[sid] = {**source, 'path': snapshot['path'], 'sha256': snapshot['sha256'],
                        'verification': '本機快照摘錄已比對；網址真實性、現行效力及適用性由 Codex 查讀，程式未連線驗證。'}
    check = data.get('legal_check')
    require(isinstance(check, dict) and check.get('status') in {'checked', 'unverified'}, '請明確提供 legal_check 查證狀態。')
    if check['status'] == 'checked':
        require(bool(catalog) and bool(jurisdiction), '宣稱完成法律來源查讀，必須提供已知法域與可核對的法律來源快照。')
        text(check.get('as_of'), '法律查讀基準日')
    else:
        text(check.get('reason'), '未完成法律查證原因')
    return catalog


def validate(data, source, base):
    for field in ['document_name', 'document_type', 'perspective']:
        text(data.get(field), field)
    jurisdiction = data.get('jurisdiction')
    require(jurisdiction is None or isinstance(jurisdiction, str), 'jurisdiction 應為文字或 null。')
    blocks = {b['location']: b['text'] for b in source['blocks']}
    checked = strings(data.get('checked_locations'), 'checked_locations')
    require(set(checked) == set(blocks), '必須完整審閱來源索引所有位置，再填 checked_locations；不能用局部引用推斷整份缺少條款。')
    topics = data.get('review_topics')
    require(isinstance(topics, list) and len(topics) == len(TOPICS), 'review_topics 必須逐項涵蓋九類條款。')
    require(all(isinstance(t, dict) for t in topics), '條款檢查必須為物件。')
    require(set(t.get('topic') for t in topics) == set(TOPICS), '條款檢查 topic 必須完整且不重複。')
    for topic in topics:
        require(topic.get('status') in {'present', 'missing', 'not_applicable'}, '條款狀態應為 present、missing、not_applicable。')
        text(topic.get('analysis'), '逐項條款分析')
        citations(topic.get('citations'), blocks, topic['status'] == 'present')
        require(topic['status'] == 'present' or not topic['citations'], '缺漏或不適用條款請以全份來源審閱範圍說明，不能引用其他文字冒充原條款。')
    laws = check_laws(data, base, jurisdiction)
    findings = data.get('findings')
    require(isinstance(findings, list), 'findings 必須為陣列。')
    ids = set()
    topic_by_id = {t['topic']: t for t in topics}
    for f in findings:
        require(isinstance(f, dict), '風險必須為物件。')
        fid = text(f.get('id'), '風險 ID')
        require(fid not in ids, '風險 ID 不可重複。')
        ids.add(fid)
        require(f.get('topic') in TOPICS, '風險必須對應條款 topic。')
        require(f.get('level') in {'高', '中', '低'}, '風險等級僅能為高、中、低。')
        require(isinstance(f.get('professional_confirmation'), bool), '請明確填寫是否需要專業法務確認。')
        for field in ['issue', 'impact', 'reason', 'action', 'draft_terms', 'revision_reason']:
            text(f.get(field), '風險 ' + field)
        strings(f.get('questions'), '待確認事項')
        missing = f.get('missing_clause', False)
        require(isinstance(missing, bool), 'missing_clause 必須為布林值。')
        require(not missing or topic_by_id[f['topic']]['status'] == 'missing', '新增缺漏條款應對應 missing 檢查項目。')
        citations(f.get('citations'), blocks, not missing)
        require(not missing or not f['citations'], '缺漏條款沒有原文，請保留空 citations 並標示新增。')
        basis = f.get('legal_basis_status')
        require(basis in {'commercial_only', 'unverified', 'supported'}, '風險需標示商業判斷、法律待查證或有來源支持。')
        refs = strings(f.get('legal_source_ids'), '法律來源 ID')
        require(all(sid in laws for sid in refs), '風險引用不存在的法律來源。')
        require(basis != 'supported' or (refs and jurisdiction and data['legal_check']['status'] == 'checked'), '法律結論需有已查讀法域及來源；未知法域只能先做商業分析與待確認事項。')
    for topic in topics:
        require(topic['status'] != 'missing' or any(f['topic'] == topic['topic'] and f.get('missing_clause') for f in findings), '缺漏關鍵條款也必須有具體風險、影響與新增修訂文字。')
    strings(data.get('questions', []), '待人工確認事項')
    return laws


def cite_text(items):
    return '\n'.join(f'[{c["location"]}] {c["quote"]}' for c in items) or '缺漏條款（新增）；已檢視全文來源索引。'


def run(input_path, output_dir):
    input_path = Path(input_path).resolve()
    data = load_json(input_path)
    source = extract_source(data.get('source'), input_path.parent)
    laws = validate(data, source, input_path.parent)
    inputs = [input_path] + ([source['path']] if source['path'] else []) + [s['path'] for s in laws.values()]
    out = fresh_output(output_dir, PRODUCTS, inputs)
    jurisdiction = data.get('jurisdiction') or '法域待確認；未假定適用台灣法律'
    label = '教學模擬資料｜' if data.get('simulation') is True else ''
    legal_limit = ('已記錄 Codex 查讀來源及日期；本產檔器未連線，不能單憑快照宣稱法規仍為最新。'
                   if data['legal_check']['status'] == 'checked' else '未完成最新法規查證：' + data['legal_check']['reason'])
    summary = '；'.join(f'{level}風險 {sum(f["level"] == level for f in data["findings"])} 項' for level in ['高', '中', '低'])
    sections = [('文件基本資訊', [f'名稱：{data["document_name"]}', f'類型：{data["document_type"]}', f'審閱立場：{data["perspective"]}', f'適用法域：{jurisdiction}', f'原件：{source["path"] or "使用者貼上的原文（見 audit.json）"}', f'SHA-256：{source["sha256"]}']),
                ('審查範圍與適用限制', [f'已逐項檢視九類條款，原文範圍共 {len(source["blocks"])} 個位置。', LIMIT, legal_limit] + source['warnings']),
                ('風險摘要', [summary, '沒有列出風險不表示文件合法或無風險；請配合逐項分析與尚待確認事項。'])]
    status_labels = {'present': '有條款', 'missing': '缺少條款', 'not_applicable': '依文件性質不適用／待確認'}
    for t in data['review_topics']:
        sections.append(('逐項審查｜' + TOPICS[t['topic']], [status_labels[t['status']], t['analysis'], cite_text(t['citations']) if t['citations'] else '依完整來源範圍判斷；無對應原條款。']))
    md = ['# ' + label + '條款修訂建議', '', LIMIT, '', f'文件：{data["document_name"]}；立場：{data["perspective"]}；法域：{jurisdiction}', legal_limit, '', '每項位置均對應本次原件擷取索引；【待確認】為未協議的條件，不表示已存在。']
    for f in data['findings']:
        body = [f'原始條款：{cite_text(f["citations"])}', f'問題：{f["issue"]}', f'商業或法律影響：{f["impact"]}', f'風險：{f["level"]}；理由：{f["reason"]}', f'處理方式：{f["action"]}', f'建議修訂文字：{f["draft_terms"]}', f'修訂理由：{f["revision_reason"]}', f'法律依據狀態：{f["legal_basis_status"]}；來源：{", ".join(f["legal_source_ids"]) or "尚無已查證法規"}', f'需專業法務確認：{"是" if f["professional_confirmation"] else "否（仍屬初步審閱）"}', '待確認：' + ('；'.join(f['questions']) or '未另列，仍需確認商業條件與附件。')]
        sections.append((f'風險 {f["id"]}｜{TOPICS[f["topic"]]}', body))
        md += ['', f'## {f["id"]}｜{TOPICS[f["topic"]]}（{f["level"]}）', ''] + body
    if not data['findings']:
        md += ['', '本次未列需修訂條款；請核對逐項審查及待確認事項，這不代表已通過正式法務審查。']
    law_lines = [legal_limit]
    for sid, law in laws.items():
        law_lines += [f'[{sid}] {law["law_name"]} {law["article"]}｜{law["authority"]}｜法域：{law["jurisdiction"]}｜查核日期：{law["checked_at"]}', law['url'], f'原文摘錄：{law["quote"]}', law['verification'], f'快照 SHA-256：{law["sha256"]}']
    questions = data.get('questions', []) + [q for f in data['findings'] for q in f['questions']]
    if not data.get('jurisdiction'):
        questions.append('請確認適用法域及履約地；在此之前無法確定法律適用或效力。')
    questions = list(dict.fromkeys(questions)) or ['請確認附件、簽署權限、實際交易背景及是否需專業法務複核。']
    sections += [('法律來源與適用限制', law_lines), ('待人工確認事項', questions), ('原始文件定位索引', [f'[{b["location"]}] {b["text"]}' for b in source['blocks']])]
    md += ['', '## 法律來源與限制', ''] + law_lines + ['', '## 待人工確認', ''] + ['- ' + q for q in questions]
    result = {**data, 'source': source, 'legal_sources': list(laws.values()), 'network_used': False,
              'semantic_review': 'Codex 提供分析；程式只核對資料結構、原文位置與快照，不以關鍵字判定法律風險。', 'limitations': [LIMIT, legal_limit] + source['warnings']}
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.legal-review-', dir=out.parent) as staging_dir:
        staging = Path(staging_dir)
        write_docx(staging / PRODUCTS[0], label + '合約審查報告', sections)
        headers = ['風險ID', '檢查項目', '原始條款／位置', '問題說明', '商業或法律影響', '風險等級', '等級理由', '處理方式', '建議修訂文字', '修訂理由', '需專業法務確認', '待確認事項', '法律依據狀態', '法律來源']
        rows = [[f['id'], TOPICS[f['topic']], cite_text(f['citations']), f['issue'], f['impact'], f['level'], f['reason'], f['action'], f['draft_terms'], f['revision_reason'], '是' if f['professional_confirmation'] else '否（初步）', '；'.join(f['questions']), f['legal_basis_status'], ', '.join(f['legal_source_ids'])] for f in data['findings']]
        write_xlsx(staging / PRODUCTS[1], headers, rows, '風險清單')
        from openpyxl import load_workbook
        wb = load_workbook(staging / PRODUCTS[1])
        overview = wb.create_sheet('審查摘要', 0)
        for row in [[label + '合約審查摘要'], ['文件', data['document_name']], ['風險摘要', summary], ['法域', jurisdiction], ['法律查證限制', legal_limit], ['判斷界線', LIMIT]]:
            overview.append([safe_cell(v) for v in row])
        overview.column_dimensions['A'].width = 32
        overview.column_dimensions['B'].width = 80
        wb.active = 0
        for name, head, rows in [('條款檢查', ['檢查項目', '狀態', '逐項分析', '原文引用'], [[TOPICS[t['topic']], status_labels[t['status']], t['analysis'], cite_text(t['citations']) if t['citations'] else '依完整文件審查'] for t in data['review_topics']]), ('來源對照', ['來源位置', '原文'], [[b['location'], b['text']] for b in source['blocks']]), ('法律來源與限制', ['項目', '內容'], [['法域', jurisdiction], ['法律查證', legal_limit], ['判斷界線', LIMIT], ['SHA-256', source['sha256']]] + [['法律來源', line] for line in law_lines]), ('待人工確認', ['待確認事項'], [[q] for q in questions])]:
            ws = wb.create_sheet(name)
            ws.append(head)
            for row in rows:
                ws.append([safe_cell(v) for v in row])
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Alignment, Font
            for row in ws:
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    cell.font = Font(name='Microsoft JhengHei', size=11, bold=cell.row == 1)
                    ws.column_dimensions[cell.column_letter].width = 48
        wb.save(staging / PRODUCTS[1])
        write_markdown(staging / PRODUCTS[2], '\n'.join(md))
        write_json(staging / PRODUCTS[3], result)
        fresh_output(out, PRODUCTS, inputs)
        out.mkdir(parents=True, exist_ok=True)
        for name in PRODUCTS:
            (staging / name).rename(out / name)
    return result


def main():
    parser = argparse.ArgumentParser(description='法務原件擷取、引用核對與審查產檔；本工具不提供自動法律判斷或 OCR。')
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--input', help='Codex 內部審查 JSON；一般使用者只需提供文件與需求。')
    mode.add_argument('--extract', help='擷取原始 DOCX/PDF/TXT/MD 及定位索引。')
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        if args.extract:
            source = extract_source({'path': str(Path(args.extract).resolve())}, Path.cwd())
            out = fresh_output(args.output, ['source-index.json'], [args.extract])
            write_json(out / 'source-index.json', source)
            print(f'已擷取 {len(source["blocks"])} 個原文位置；請由 Codex 逐項語意審查。')
        else:
            result = run(args.input, args.output)
            print(f'已核對原文並產出三份審查成果；{len(result["findings"])} 項風險。法律查讀與判斷限制見報告。')
        return 0
    except (ValueError, OSError, TypeError, KeyError, ImportError) as exc:
        print('無法完成：' + str(exc), file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
