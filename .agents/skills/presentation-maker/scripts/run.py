"""把 Codex 已依來源整理的 JSON 製成可編輯 PPTX 與講者備註。"""
import argparse
import math
import sys
import unicodedata
from pathlib import Path

from _office import load_json, require, write_json, write_markdown, protect_input


def units(text):
    return sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1 for c in text)


def text(value, label, required=True):
    require(isinstance(value, str) and (bool(value.strip()) or not required), f'{label} 必須為文字且不可空白。')
    return value.strip()


def validate(data):
    title = text(data.get('title'), '簡報標題')
    source = text(data.get('source'), '來源 source')
    require(units(title) <= 72 and '\n' not in title, '簡報標題過長或含換行，請縮短至 36 個全形字以內。')
    slides = data.get('slides')
    require(isinstance(slides, list) and 1 <= len(slides) <= 60, 'slides 應為 1 至 60 頁的陣列；封面會另加 1 頁。')
    cleaned = []
    for index, slide in enumerate(slides, 1):
        require(isinstance(slide, dict), f'第 {index} 頁必須是物件。')
        slide_title = text(slide.get('title'), f'第 {index} 頁標題')
        require(units(slide_title) <= 54 and '\n' not in slide_title, f'第 {index} 頁標題過長，請縮短至 27 個全形字以內。')
        bullets = slide.get('bullets')
        require(isinstance(bullets, list) and 1 <= len(bullets) <= 5, f'第 {index} 頁須有 1 至 5 個重點，過多請拆成數頁。')
        bullets = [text(v, f'第 {index} 頁重點') for v in bullets]
        require(all(units(v) <= 84 and '\n' not in v for v in bullets), f'第 {index} 頁重點過長或含換行；每點最多 42 個全形字，請由 Codex 依原意拆頁，程式不會截斷。')
        require(sum(units(v) for v in bullets) <= 260, f'第 {index} 頁文字量過多，請保留重點並拆頁。')
        # 保守估算斷行及段落間距，避免總字量合格但多個長段落碰到頁尾。
        wrapped_lines = sum(math.ceil((sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1.6 for c in v) + 8) / 60) for v in bullets)
        estimated_height = wrapped_lines * 32 + (len(bullets) - 1) * 17
        require(estimated_height <= 260, f'第 {index} 頁預估斷行後文字高度過高，請減少重點或拆頁。')
        notes = slide.get('notes')
        if notes is None or notes == '':
            notes = '待確認：尚未提供講者備註，請依原始資料補寫。'
        else:
            notes = text(notes, f'第 {index} 頁講者備註')
        require(len(notes) <= 12000, f'第 {index} 頁備註過長，請拆頁。')
        slide_source = text(slide.get('source') or source, f'第 {index} 頁來源')
        cleaned.append({'title': slide_title, 'bullets': bullets, 'notes': notes, 'source': slide_source})
    require(isinstance(data.get('simulation', False), bool), 'simulation 必須是 true 或 false。')
    return title, source, cleaned


def add_text(slide, value, x, y, w, h, size, color, bold=False):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.name = 'Microsoft JhengHei'
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    return shape


def generate(args):
    protect_input(args.input, args.output)
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    data = load_json(args.input)
    title, source, slides = validate(data)
    out = Path(args.output)
    require(Path(args.input).resolve() != (out / 'presentation.pptx').resolve(), '輸出不可覆寫來源檔案。')
    out.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = title
    deck.core_properties.author = 'AI 辦公室工具包'
    deck.core_properties.subject = source
    simulation = bool(data.get('simulation', False))
    label = '模擬教學資料' if simulation else '依使用者提供資料製作'
    notes_markdown = [f'# {title}｜講者備註', '', f'來源：{source}', f'資料標示：{label}', '',
                      '待人工檢查：請在 PowerPoint 開啟並檢視每頁畫面。程式只完成結構與字量檢查，未宣稱實際渲染驗收。', '']
    cover = deck.slides.add_slide(deck.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor.from_string('123E3B')
    add_text(cover, label, .8, .65, 11.5, .6, 17, 'B6D7D1')
    add_text(cover, title, .8, 2.25, 11.5, 2.1, 44, 'FFFFFF', True)
    subtitle = data.get('subtitle')
    if subtitle:
        subtitle = text(subtitle, '副標題')
        require(units(subtitle) <= 90 and '\n' not in subtitle, '副標題過長，請縮短至 45 個全形字以內。')
        add_text(cover, subtitle, .8, 4.8, 11.5, 1.0, 22, 'DFEEEB')
    add_text(cover, '來源與資料日期見講者備註', .8, 6.55, 11.5, .4, 12, 'B6D7D1')
    cover.notes_slide.notes_text_frame.text = f'{label}\n來源：{source}\n資料日期：{data.get("date") or "未提供，待確認"}'
    notes_markdown.extend(['## 封面', '', cover.notes_slide.notes_text_frame.text, ''])
    for index, item in enumerate(slides, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('F8FAF9')
        add_text(slide, f'{index:02d}  {label}', .8, .45, 11.7, .4, 12, '537571')
        add_text(slide, item['title'], .8, 1.0, 11.7, 1.2, 32, '123E3B', True)
        shape = add_text(slide, '', .9, 2.5, 11.4, 3.9, 24, '203C39')
        for i, bullet in enumerate(item['bullets']):
            paragraph = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            paragraph.text = f'{i + 1}.  {bullet}'
            paragraph.font.name = 'Microsoft JhengHei'
            paragraph.font.size = Pt(24)
            paragraph.font.color.rgb = RGBColor.from_string('203C39')
            paragraph.space_after = Pt(17)
            paragraph.line_spacing = 1.2
        add_text(slide, f'來源與核對事項見講者備註       {index + 1} / {len(slides) + 1}', .8, 6.85, 11.7, .35, 11, '537571')
        notes = f'{item["notes"]}\n\n來源：{item["source"]}\n資料日期：{data.get("date") or "未提供，待確認"}\n資料標示：{label}'
        slide.notes_slide.notes_text_frame.text = notes
        notes_markdown.extend([f'## {index + 1}. {item["title"]}', '', notes, ''])
    checks = []
    for index, slide in enumerate(deck.slides, 1):
        failures = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width or shape.top + shape.height > deck.slide_height:
                failures.append(shape.shape_id)
        checks.append({'slide': index, 'shape_count': len(slide.shapes), 'out_of_bounds_shapes': failures,
                       'has_speaker_notes': bool(slide.notes_slide.notes_text_frame.text.strip())})
    require(all(not c['out_of_bounds_shapes'] and c['has_speaker_notes'] for c in checks), '頁面結構檢查失敗，請檢查輸入字量及版面。')
    deck.save(out / 'presentation.pptx')
    # 寫檔後重新開啟，確保 ZIP / OOXML 可解析且 notes 實際存在。
    reopened = Presentation(out / 'presentation.pptx')
    require(len(reopened.slides) == len(slides) + 1, 'PPTX 重新開啟後頁數不符。')
    require(all(s.notes_slide.notes_text_frame.text.strip() for s in reopened.slides), 'PPTX 備註未完整保存。')
    write_json(out / 'layout-check.json', {'structure_passed': True, 'visual_review_completed': False,
               'slide_count': len(reopened.slides), 'checks': checks,
               'text_rule': '32pt 頁標、24pt 內文；每點最多84寬度單位、每頁最多260單位／5點，另以保守斷行與段距預算限制預估高度260pt。',
               'remaining_review': '需要 PowerPoint 或可用渲染工具檢查實際字型替代、斷行與文字溢出。'})
    write_markdown(out / 'speaker-notes.md', '\n'.join(notes_markdown))
    print('已產生 PPTX、講者備註與結構檢查：' + str(out.resolve()))


def main():
    parser = argparse.ArgumentParser(description='依有來源的簡報 JSON 產生實際 PPTX；長文字會明確拒絕。')
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        generate(args)
        return 0
    except ImportError:
        print('錯誤：缺少 python-pptx，請安裝工具包 requirements.txt。', file=sys.stderr)
    except Exception as exc:
        print('錯誤：無法完成簡報。' + str(exc), file=sys.stderr)
    return 1


if __name__ == '__main__':
    sys.exit(main())
