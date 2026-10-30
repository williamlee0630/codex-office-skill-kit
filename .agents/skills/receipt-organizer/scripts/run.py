"""收據資料整理；影像只建立辨識證據，不冒稱 OCR。"""
import argparse
import csv
import hashlib
import importlib.util
import math
import re
import shutil
import sys
import unicodedata
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from _office import load_json, require, safe_cell, write_json, write_markdown, write_xlsx, protect_input


def empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def tool_check():
    return {'pypdf': importlib.util.find_spec('pypdf') is not None,
            'Pillow': importlib.util.find_spec('PIL') is not None,
            'tesseract_path': shutil.which('tesseract'),
            'pdftoppm_path': shutil.which('pdftoppm'),
            'note': '有工具不代表已辨識；本腳本不自動執行 OCR。文字型 PDF 使用 pypdf 擷取。'}


def preserve(path, out):
    original = out / 'originals' / path.name
    original.parent.mkdir(parents=True, exist_ok=True)
    if original.resolve() != path.resolve():
        # 同名不同內容不得覆寫，避免混淆原始證據。
        if original.exists():
            require(original.read_bytes() == path.read_bytes(), 'originals 已有同名但不同內容的檔案，請改用新的輸出資料夾。')
        else:
            shutil.copy2(path, original)
    return {'input_path': str(path.resolve()), 'preserved_path': str(original.resolve()),
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def prepare_raw(path, out, tools, simulation=False):
    require(not (out / 'receipts.xlsx').exists(), '輸出資料夾已有報帳 Excel；請為影像辨識使用新的輸出資料夾，以免誤用舊報表。')
    evidence = preserve(path, out)
    pages, image_info = [], None
    if path.suffix.lower() == '.pdf':
        require(tools['pypdf'], '缺少 pypdf，請安裝 requirements.txt 或提供人工轉錄 JSON／CSV。')
        from pypdf import PdfReader
        reader = PdfReader(path)
        require(not reader.is_encrypted, 'PDF 已加密，請提供可讀取且已授權的副本。')
        for number, page in enumerate(reader.pages, 1):
            pages.append({'page': number, 'text': page.extract_text() or ''})
    else:
        require(tools['Pillow'], '缺少 Pillow，請安裝 requirements.txt 或提供人工轉錄 JSON／CSV。')
        from PIL import Image
        with Image.open(path) as image:
            image_info = {'format': image.format, 'width': image.width, 'height': image.height}
            image.verify()
    extracted = '\n\n'.join(f'第 {p["page"]} 頁\n{p["text"]}' for p in pages)
    has_extracted_text = any(p['text'].strip() for p in pages)
    evidence.update(status='needs_review', simulation=simulation, ocr_performed=False, extraction_tools=tools,
                    extracted_text=extracted, pages=pages, image_info=image_info,
                    instruction='請由 Codex 可用的視覺工具或人工逐欄確認後填入 transcription-template.json。未辨識值保留 null，並註明頁碼／區域與原文證據。')
    write_json(out / 'extraction-evidence.json', evidence)
    write_json(out / 'transcription-template.json', {'source': path.name, 'simulation': simulation, 'receipts': [
        {'id': 'R1', 'date': None, 'vendor': None, 'items': None, 'amount': None, 'currency': None,
         'category': None, 'evidence': f'待人工確認：{path.name}，請補頁碼／區域與辨識原文'}]})
    write_markdown(out / 'review-required.md', f'# 收據待核對\n\n資料標示：{"模擬教學資料" if simulation else "依使用者提供來源整理"}\n\n來源：{path.name}\n\n原始檔已保留於 originals，SHA-256 與抽取工具清單見 extraction-evidence.json。\n\n'
                   + ('已擷取 PDF 文字，但欄位仍需 Codex 根據頁面確認。' if has_extracted_text else '此來源未擷取到文字，請用可用視覺工具或人工轉錄。')
                   + '\n\n本次未執行 OCR，尚未產生報帳 Excel。請填好 transcription-template.json，保留未知欄位，再以該 JSON 為輸入產生報表。\n')
    print('原始檔與擷取證據已保留，欄位待人工確認：' + str(out.resolve()))


def money(value, index):
    if empty(value):
        return None
    require(not isinstance(value, bool), f'第 {index} 筆金額不能是布林值。')
    try:
        amount = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f'第 {index} 筆金額格式錯誤，未知金額請填 null，不可猜測。') from exc
    require(amount.is_finite(), f'第 {index} 筆金額必須是有限數值，不能為 NaN 或無限值。')
    require(amount >= 0, f'第 {index} 筆為負數，退款須另由人工確認，請勿混入一般收據。')
    require(amount.as_tuple().exponent >= -2, f'第 {index} 筆金額超過兩位小數，請依原始收據核對，不自動四捨五入。')
    require(amount <= Decimal('999999999999'), f'第 {index} 筆金額超過本工具可安全輸出的範圍。')
    return amount.quantize(Decimal('0.01'))


def load_records(path):
    if path.suffix.lower() == '.json':
        data = load_json(path)
    elif path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            require(reader.fieldnames and len(reader.fieldnames) == len(set(reader.fieldnames)), 'CSV 需要不重複的欄位名稱。')
            records = list(reader)
        require(all(None not in record for record in records), 'CSV 存在多餘欄位，請檢查分隔符號。')
        data = {'source': path.name, 'simulation': False, 'receipts': records}
    else:
        raise ValueError('可解析資料僅支援 JSON 或 CSV；影像／PDF 須先完成欄位確認。')
    require(isinstance(data.get('source'), str) and data['source'].strip(), '需要 source 來源名稱。')
    require(isinstance(data.get('receipts'), list) and bool(data['receipts']), 'receipts 必須包含至少一筆收據。')
    require(isinstance(data.get('simulation', False), bool), 'simulation 必須是 true 或 false。')
    return data


def normalize(data):
    records, seen, ids = [], {}, set()
    totals = defaultdict(lambda: {'all': Decimal(0), 'excluding_suspected_duplicates': Decimal(0), 'count': 0})
    category_totals = defaultdict(Decimal)
    for index, raw in enumerate(data['receipts'], 1):
        require(isinstance(raw, dict), f'第 {index} 筆收據必須是物件。')
        entry = dict(raw)
        for field in ('id', 'date', 'vendor', 'items', 'currency', 'category', 'evidence'):
            require(empty(entry.get(field)) or isinstance(entry.get(field), str), f'第 {index} 筆 {field} 必須為文字或 null。')
            if empty(entry.get(field)):
                entry[field] = None
        entry['id'] = entry['id'] or f'row-{index}'
        require(entry['id'] not in ids, f'收據 id 重複：{entry["id"]}，請提供唯一識別碼。')
        ids.add(entry['id'])
        amount = money(entry.get('amount'), index)
        entry['amount'] = str(amount) if amount is not None else None
        entry['missing_fields'] = [key for key in ('date', 'vendor', 'items', 'amount', 'currency', 'category', 'evidence') if empty(entry.get(key))]
        entry['invalid_fields'] = []
        if entry['date']:
            try:
                require(date.fromisoformat(entry['date']).isoformat() == entry['date'], '日期格式不符')
            except ValueError:
                entry['invalid_fields'].append('date')
        if entry['currency']:
            entry['currency'] = entry['currency'].strip().upper()
            if not re.fullmatch('[A-Z]{3}', entry['currency']):
                entry['invalid_fields'].append('currency')
        entry['category'] = entry['category'] or '待人工確認'
        entry['duplicate_of'] = None
        # 只以來源可核對的精確鍵標記疑似重複；不刪除、不自動認定為錯誤。
        if amount is not None and entry['date'] and entry['vendor'] and entry['currency'] and not entry['invalid_fields']:
            key = (entry['date'], entry['vendor'].strip(), entry.get('items'), entry['amount'], entry['currency'])
            if key in seen:
                entry['duplicate_of'] = seen[key]
            else:
                seen[key] = entry['id']
        entry['needs_review'] = bool(entry['missing_fields'] or entry['invalid_fields'] or entry['duplicate_of'] or entry['category'] == '待人工確認')
        if amount is not None and entry['currency'] and 'currency' not in entry['invalid_fields']:
            currency = entry['currency']
            totals[currency]['all'] += amount
            totals[currency]['count'] += 1
            if not entry['duplicate_of']:
                totals[currency]['excluding_suspected_duplicates'] += amount
            category_totals[(currency, entry['category'])] += amount
        records.append(entry)
    return records, {key: {k: (f'{v:.2f}' if isinstance(v, Decimal) else v) for k, v in value.items()} for key, value in sorted(totals.items())}, [
        {'currency': key[0], 'category': key[1], 'total_including_duplicates': f'{value:.2f}'} for key, value in sorted(category_totals.items())]


def workbook(path, records, totals, categories, source):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    headers = ['識別碼', '日期', '店家', '品項', '金額', '幣別', '類別', '疑似重複於', '待人工確認', '缺漏欄位', '格式待核對', '辨識證據']
    rows = [[r['id'], r.get('date'), r.get('vendor'), r.get('items'), float(r['amount']) if r['amount'] is not None else None,
             r.get('currency'), r['category'], r['duplicate_of'], '是' if r['needs_review'] else '否',
             ', '.join(r['missing_fields']), ', '.join(r['invalid_fields']), r.get('evidence')] for r in records]
    write_xlsx(path, headers, rows, sheet_name='報帳明細')
    wb = load_workbook(path)
    for cell in wb['報帳明細']['E'][1:]:
        cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
    ws = wb.create_sheet('各幣統計')
    ws.append(['幣別', '有效金額筆數', '含疑似重複合計', '暫排除疑似重複合計'])
    for currency, info in totals.items():
        ws.append([currency, info['count'], float(info['all']), float(info['excluding_suspected_duplicates'])])
    ws.append([])
    ws.append(['不同幣別分開合計，未換匯。所有收據均保留；疑似重複仍需人工判斷。'])
    ws.append(['缺少金額或有效幣別的記錄不納入合計，請見報帳明細。'])
    ws.append(['資料來源', safe_cell(source)])
    ws.append(['結果為靜態計算，修改來源後請重新執行。'])
    note_start = len(totals) + 3
    ws.merge_cells(start_row=note_start, start_column=1, end_row=note_start, end_column=4)
    ws.merge_cells(start_row=note_start + 1, start_column=1, end_row=note_start + 1, end_column=4)
    ws.merge_cells(start_row=note_start + 2, start_column=2, end_row=note_start + 2, end_column=4)
    ws.merge_cells(start_row=note_start + 3, start_column=1, end_row=note_start + 3, end_column=4)
    cat = wb.create_sheet('分類統計')
    cat.append(['幣別', '使用者／Codex 確認類別', '含疑似重複合計'])
    for item in categories:
        cat.append([item['currency'], safe_cell(item['category']), float(item['total_including_duplicates'])])
    wb._sheets = [wb['各幣統計'], wb['分類統計'], wb['報帳明細']]
    from openpyxl.utils import get_column_letter
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{len(records) + 1}' if sheet.title == '報帳明細' else f'A1:{get_column_letter(sheet.max_column)}{len(totals) + 1 if sheet.title == "各幣統計" else len(categories) + 1}'
        for cell in sheet[1]:
            cell.font = Font(name='Microsoft JhengHei', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='164E4A')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Microsoft JhengHei', size=11)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 24
        if sheet.title == '報帳明細':
            sheet.column_dimensions['L'].width = 55
        elif sheet.title == '各幣統計':
            sheet.column_dimensions['C'].width = sheet.column_dimensions['D'].width = 30
        for row in sheet.iter_rows(min_row=2):
            line_count = 1
            for cell in row:
                if cell.value is None:
                    continue
                width = sheet.column_dimensions[get_column_letter(cell.column)].width or 24
                for merged in sheet.merged_cells.ranges:
                    if cell.coordinate in merged:
                        width = sum(sheet.column_dimensions[get_column_letter(c)].width or 24 for c in range(merged.min_col, merged.max_col + 1))
                        break
                units = sum(2 if unicodedata.east_asian_width(c) in {'W', 'F'} else 1 for c in str(cell.value))
                line_count = max(line_count, math.ceil(units / max(1, width - 2)))
            sheet.row_dimensions[row[0].row].height = min(409, max(35, line_count * 15 + 8))
    wb.save(path)


def organize(args):
    protect_input(args.input, args.output)
    path, out = Path(args.input), Path(args.output)
    require(path.is_file(), '找不到輸入收據。')
    require(path.resolve() not in {(out / 'receipts-summary.json').resolve(), (out / 'transcription-template.json').resolve()}, '請使用不同輸出資料夾，避免覆寫輸入資料。')
    tools = tool_check()
    if path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.tif', '.tiff', '.bmp', '.pdf'}:
        out.mkdir(parents=True, exist_ok=True)
        prepare_raw(path, out, tools, args.simulation)
        return
    data = load_records(path)
    if args.simulation:
        data['simulation'] = True
    records, totals, categories = normalize(data)
    out.mkdir(parents=True, exist_ok=True)
    original = preserve(path, out)
    result = {'source': data['source'], 'simulation': data.get('simulation', False), 'original': original,
              'extraction_tools': tools, 'ocr_performed': False, 'receipts': records, 'totals': totals,
              'category_totals': categories, 'review_count': sum(r['needs_review'] for r in records),
              'limitations': ['分類依輸入值，未提供類別即待人工確認。', '不同幣別不混加、不換匯。', '疑似重複均保留，未當作已確認的重複報帳。', '靜態結果，更新來源後重新執行。']}
    write_json(out / 'receipts-summary.json', result)
    workbook(out / 'receipts.xlsx', records, totals, categories, data['source'])
    label = '模擬教學資料' if data.get('simulation') else '依使用者提供來源整理'
    lines = ['# 收據整理報告', '', f'資料標示：{label}。來源：{data["source"]}',
             f'共 {len(records)} 筆，待人工確認 {result["review_count"]} 筆。未執行 OCR。', '',
             '| 幣別 | 含疑似重複合計 | 暫排除疑似重複合計 |', '| --- | ---: | ---: |']
    for currency, info in totals.items():
        lines.append(f'| {currency} | {info["all"]} | {info["excluding_suspected_duplicates"]} |')
    lines.extend(['', '不同幣別分別合計，未換匯。沒有有效金額或幣別的收據不納入合計。', '', '## 待人工確認', ''])
    for r in records:
        if r['needs_review']:
            lines.append(f'- {r["id"]}：缺漏 {", ".join(r["missing_fields"]) or "無"}；格式待核對 {", ".join(r["invalid_fields"]) or "無"}；疑似重複於 {r["duplicate_of"] or "無"}。證據：{r.get("evidence") or "未提供"}')
    lines.extend(['', '原始檔與 SHA-256 已保留。類別依使用者／Codex 確認的輸入，腳本不依店家猜測用途。',
                  '請逐筆核對原始影像／PDF、日期、品項、金額與幣別；疑似重複不代表已確認錯誤。',
                  'Excel 與 JSON 為 Python 靜態計算結果，修改來源後請重新執行。'])
    write_markdown(out / 'receipts-report.md', '\n'.join(lines))
    print('已產生報帳 Excel、各幣合計與人工確認清單：' + str(out.resolve()))


def main():
    parser = argparse.ArgumentParser(description='整理 JSON／CSV 收據資料，或保留 PDF／影像與擷取證據。')
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--simulation', action='store_true', help='明確標記本次輸入為模擬教學資料；不根據檔名猜測。')
    args = parser.parse_args()
    try:
        organize(args)
        return 0
    except ImportError:
        print('錯誤：缺少格式套件，請安裝工具包 requirements.txt。', file=sys.stderr)
    except Exception as exc:
        print('錯誤：無法完成收據整理。' + str(exc), file=sys.stderr)
    return 1


if __name__ == '__main__':
    sys.exit(main())
