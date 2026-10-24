"""第三階段：真實產檔、例外與安全行為測試（測試資料皆為模擬）。"""
import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


class ArtifactTests(unittest.TestCase):
    def test_presentation_rejects_wrapped_lines_that_reach_footer(self):
        data = {'title': '密集文字測試', 'source': '模擬來源', 'slides': [
            {'title': '密集頁', 'bullets': ['測' * 42, '試' * 42, '字' * 42, '四點', '五點']}]}
        self.run_skill('presentation-maker', self.write_json(data), ok=False)
        self.assertFalse((self.out / 'presentation.pptx').exists())
    def test_presentation_preserves_source_named_like_output(self):
        path = self.work / 'layout-check.json'
        original = json.dumps({'title': '模擬簡報', 'source': '模擬來源', 'slides': [{'title': '測試', 'bullets': ['內容']}]}, ensure_ascii=False)
        path.write_text(original, encoding='utf-8')
        result = subprocess.run([sys.executable, '-X', 'utf8', str(ROOT / '.agents/skills/presentation-maker/scripts/run.py'),
                                 '--input', str(path), '--output', str(self.work)], capture_output=True, text=True, encoding='utf-8')
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(path.read_text(encoding='utf-8'), original)

    def test_data_markdown_table_is_not_split_by_outlier_paragraph(self):
        self.run_skill('data-analyst', ROOT / 'examples/data-analyst/sales.csv')
        text = (self.out / 'analysis.md').read_text(encoding='utf-8')
        section = text.split('## 描述統計', 1)[1].split('## 圖表與比較', 1)[0]
        lines = section.splitlines()
        table_lines = [i for i, line in enumerate(lines) if line.startswith('|')]
        self.assertEqual(table_lines, list(range(min(table_lines), max(table_lines)+1)))
    def setUp(self):
        temp_root = ROOT / 'tests' / '.tmp'
        temp_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=temp_root)
        self.work = Path(self.temp.name)
        self.out = self.work / "output"

    def tearDown(self):
        self.temp.cleanup()

    def run_skill(self, name, path, *args, ok=True):
        result = subprocess.run([sys.executable, str(ROOT / '.agents/skills' / name / 'scripts/run.py'),
                                 '--input', str(path), '--output', str(self.out), *args],
                                capture_output=True, text=True, encoding='utf-8')
        if ok:
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        else:
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('錯誤', result.stderr)
        return result

    def write_json(self, data):
        path = self.work / 'input.json'
        path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        return path

    def write_csv(self, text):
        path = self.work / 'input.csv'
        path.write_text(text, encoding='utf-8-sig')
        return path

    def test_data_statistics_groups_trend_chart_and_safe_workbook(self):
        from openpyxl import load_workbook
        from PIL import Image
        path = self.write_csv('date,category,amount,note\n2026-10-01,A,100,=1+1\n2026-10-02,A,200,一般\n2026-10-03,B,300,一般\n2026-10-04,B,,缺漏\n')
        self.run_skill('data-analyst', path, '--group-by', 'category', '--value', 'amount', '--date-column', 'date')
        data = json.loads((self.out / 'analysis.json').read_text(encoding='utf-8'))
        self.assertEqual(data['columns']['amount']['sum'], 600)
        self.assertEqual(data['columns']['amount']['mean'], 200)
        self.assertEqual(data['columns']['amount']['missing_count'], 1)
        self.assertEqual(data['groups']['A']['sum'], 300)
        self.assertEqual(data['trend'][0]['date'], '2026-10-01')
        self.assertEqual(data['trend'][-1]['sum'], 300)
        wb = load_workbook(self.out / 'analysis.xlsx')
        self.assertNotEqual(wb['原始資料']['D2'].data_type, 'f')
        self.assertIn('因果', (self.out / 'analysis.md').read_text(encoding='utf-8'))
        with Image.open(self.out / 'chart.png') as image:
            self.assertGreaterEqual(image.width, 800)

    def test_data_xlsx_single_point_and_incomplete(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(['category', 'amount'])
        wb.active.append(['模擬', 42])
        wb.active.append(['', None])
        path = self.work / 'input.xlsx'
        wb.save(path)
        self.run_skill('data-analyst', path)
        data = json.loads((self.out / 'analysis.json').read_text(encoding='utf-8'))
        self.assertEqual(data['row_count'], 1)
        self.assertIsNone(data['columns']['amount']['sample_stdev'])

    def test_data_empty_bad_shape_and_unsupported(self):
        self.run_skill('data-analyst', self.write_csv('amount,category\n'), ok=False)
        self.run_skill('data-analyst', self.write_csv('amount,amount\n1,2\n'), ok=False)
        path = self.work / 'input.txt'
        path.write_text('模擬錯誤資料', encoding='utf-8')
        self.run_skill('data-analyst', path, ok=False)

    def test_data_outliers_and_invalid_values_preserved(self):
        path = self.write_csv('amount\n10\n10\n11\n12\n1000\n文字\n')
        self.run_skill('data-analyst', path)
        data = json.loads((self.out / 'analysis.json').read_text(encoding='utf-8'))
        self.assertEqual(data['columns']['amount']['invalid_count'], 1)
        self.assertEqual(data['columns']['amount']['outlier_rows'], [6])

    def test_data_rejects_stale_output(self):
        self.out.mkdir()
        (self.out / 'chart.png').write_bytes(b'old-chart')
        self.run_skill('data-analyst', self.write_csv('source\n模擬文字資料\n'), ok=False)
        self.assertEqual((self.out / 'chart.png').read_bytes(), b'old-chart')

    def test_presentation_real_pptx_notes_and_layout(self):
        from pptx import Presentation
        data = {'title': '模擬季度報告', 'source': 'source.md', 'simulation': True,
                'slides': [{'title': '本季進度', 'bullets': ['已完成模擬流程演練', '未確認事項保留'], 'notes': '來源：source.md 第 1 段。'}]}
        self.run_skill('presentation-maker', self.write_json(data))
        deck = Presentation(self.out / 'presentation.pptx')
        self.assertEqual(len(deck.slides), 2)
        self.assertIn('來源：source.md', deck.slides[1].notes_slide.notes_text_frame.text)
        layout = json.loads((self.out / 'layout-check.json').read_text(encoding='utf-8'))
        self.assertTrue(layout['structure_passed'])
        self.assertFalse(layout['visual_review_completed'])
        self.assertIn('待人工', (self.out / 'speaker-notes.md').read_text(encoding='utf-8'))

    def test_presentation_missing_fields_and_long_text_are_explicit(self):
        self.run_skill('presentation-maker', self.write_json({'title': '模擬', 'source': '模擬.md', 'slides': [{'title': '進度', 'bullets': ['資料待確認']}]}))
        self.assertIn('待確認', (self.out / 'speaker-notes.md').read_text(encoding='utf-8'))
        self.run_skill('presentation-maker', self.write_json({'title': '模擬', 'source': '模擬.md', 'slides': [{'title': '過長', 'bullets': ['長' * 500]}]}), ok=False)
        self.run_skill('presentation-maker', self.write_json({'title': '無來源', 'slides': []}), ok=False)

    def test_presentation_rejects_raw_and_keeps_input(self):
        path = self.work / 'source.md'
        path.write_text('# 模擬報告\n尚待 Codex 整理', encoding='utf-8')
        self.run_skill('presentation-maker', path, ok=False)
        self.assertIn('尚待 Codex', path.read_text(encoding='utf-8'))

    def receipt_data(self):
        return {'simulation': True, 'source': '模擬報帳清單', 'receipts': [
            {'id': 'R1', 'date': '2026-10-01', 'vendor': '模擬餐館', 'items': '午餐', 'amount': '120', 'currency': 'TWD', 'category': '餐飲', 'evidence': '模擬轉錄第 1 筆'},
            {'id': 'R2', 'date': '2026-10-01', 'vendor': '模擬餐館', 'items': '午餐', 'amount': '120', 'currency': 'TWD', 'category': '餐飲', 'evidence': '模擬轉錄第 2 筆'},
            {'id': 'R3', 'date': '2026-10-02', 'vendor': '=1+1', 'items': '模擬票', 'amount': '10', 'currency': 'USD', 'category': '交通', 'evidence': '模擬轉錄第 3 筆'},
            {'id': 'R4', 'date': None, 'vendor': None, 'amount': None, 'currency': None, 'evidence': '字跡不清'}]}

    def test_receipt_currency_duplicates_missing_evidence_and_safe_xlsx(self):
        from openpyxl import load_workbook
        self.run_skill('receipt-organizer', self.write_json(self.receipt_data()))
        data = json.loads((self.out / 'receipts-summary.json').read_text(encoding='utf-8'))
        self.assertEqual(data['totals']['TWD']['all'], '240.00')
        self.assertEqual(data['totals']['TWD']['excluding_suspected_duplicates'], '120.00')
        self.assertEqual(data['totals']['USD']['all'], '10.00')
        self.assertEqual(data['receipts'][1]['duplicate_of'], 'R1')
        self.assertIn('amount', data['receipts'][3]['missing_fields'])
        self.assertEqual(len(data['receipts']), 4)
        wb = load_workbook(self.out / 'receipts.xlsx')
        self.assertNotEqual(wb['報帳明細']['C4'].data_type, 'f')
        self.assertIn('人工確認', (self.out / 'receipts-report.md').read_text(encoding='utf-8'))

    def test_receipt_csv_and_invalid_money(self):
        path = self.write_csv('id,date,vendor,items,amount,currency,category,evidence\nR1,2026-10-01,模擬店,紙張,20,TWD,辦公用品,第1行\n')
        self.run_skill('receipt-organizer', path)
        data = self.receipt_data()
        data['receipts'][0]['amount'] = 'NaN'
        self.run_skill('receipt-organizer', self.write_json(data), ok=False)
        self.run_skill('receipt-organizer', self.write_json({'source': '模擬', 'receipts': []}), ok=False)

    def test_receipt_unknown_currency_and_missing_evidence_not_silently_approved(self):
        data = self.receipt_data()
        data['receipts'][0]['currency'] = '未知'
        data['receipts'][2]['evidence'] = None
        self.run_skill('receipt-organizer', self.write_json(data))
        result = json.loads((self.out / 'receipts-summary.json').read_text(encoding='utf-8'))
        self.assertIn('currency', result['receipts'][0]['invalid_fields'])
        self.assertEqual(result['totals']['TWD']['all'], '120.00')
        self.assertIn('evidence', result['receipts'][2]['missing_fields'])

    def test_receipt_simulation_flag_for_raw_input(self):
        from PIL import Image
        path = self.work / 'receipt.png'
        Image.new('RGB', (200, 300), 'white').save(path)
        self.run_skill('receipt-organizer', path, '--simulation')
        data = json.loads((self.out / 'transcription-template.json').read_text(encoding='utf-8'))
        self.assertTrue(data['simulation'])

    def test_receipt_long_evidence_gets_readable_row_and_merged_notes(self):
        from openpyxl import load_workbook
        data = self.receipt_data()
        data['receipts'][0]['evidence'] = '模擬辨識證據' * 30
        self.run_skill('receipt-organizer', self.write_json(data))
        wb = load_workbook(self.out / 'receipts.xlsx')
        self.assertGreater(wb['報帳明細'].row_dimensions[2].height, 35)
        self.assertTrue(wb['各幣統計'].merged_cells.ranges)

    def test_receipt_image_and_pdf_are_preserved_and_not_fake_ocr(self):
        from PIL import Image
        from reportlab.pdfgen import canvas
        path = self.work / '模擬收據.png'
        Image.new('RGB', (300, 400), 'white').save(path)
        self.run_skill('receipt-organizer', path)
        data = json.loads((self.out / 'extraction-evidence.json').read_text(encoding='utf-8'))
        self.assertEqual(data['status'], 'needs_review')
        self.assertFalse(data['ocr_performed'])
        self.assertTrue((self.out / 'originals' / path.name).exists())
        self.assertEqual((self.out / 'originals' / path.name).read_bytes(), path.read_bytes())
        self.assertFalse((self.out / 'receipts.xlsx').exists())
        pdf = self.work / '模擬收據.pdf'
        c = canvas.Canvas(str(pdf)); c.drawString(40, 700, 'SIMULATED RECEIPT TWD 120'); c.save()
        self.run_skill('receipt-organizer', pdf)
        data = json.loads((self.out / 'extraction-evidence.json').read_text(encoding='utf-8'))
        self.assertIn('SIMULATED', data['extracted_text'])
        self.assertFalse(data['ocr_performed'])

    def test_each_skill_runs_after_standalone_copy(self):
        for name, source in [('data-analyst', 'sales.csv'), ('presentation-maker', 'deck.json'), ('receipt-organizer', 'receipts.json')]:
            with self.subTest(skill=name):
                copied = self.work / name
                shutil.copytree(ROOT / '.agents/skills' / name, copied)
                result = subprocess.run([sys.executable, str(copied / 'scripts/run.py'), '--input', str(ROOT / 'examples' / name / source), '--output', str(self.work / (name + '-out'))], capture_output=True, text=True, encoding='utf-8')
                self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == '__main__':
    unittest.main()
