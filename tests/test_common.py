import importlib.util
import json
from pathlib import Path
import tempfile
import subprocess
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
TEMP = ROOT / 'tests/.tmp'
TEMP.mkdir(parents=True, exist_ok=True)


class CommonTests(unittest.TestCase):
    def test_extract_cli_refuses_to_replace_original_document(self):
        from docx import Document
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            path = Path(temp) / 'source.docx'
            document = Document()
            document.add_paragraph('原始模擬文件必須完整保留。')
            document.save(path)
            original = path.read_bytes()
            result = subprocess.run([sys.executable, '-X', 'utf8', str(ROOT / 'tools/office_common.py'),
                                     'extract', str(path), '--output', str(path)], capture_output=True, text=True, encoding='utf-8')
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(path.read_bytes(), original)

    def setUp(self):
        self.assertTrue((ROOT / 'tools/office_common.py').is_file(), '尚未實作共用產檔工具')
        spec = importlib.util.spec_from_file_location('office_common', ROOT / 'tools/office_common.py')
        self.common = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.common)

    def test_json_rejects_wrong_root_and_nonfinite(self):
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'input.json'
            p.write_text('[]', encoding='utf-8')
            with self.assertRaises(ValueError):
                self.common.load_json(p)
            p.write_text('{"amount": NaN}', encoding='utf-8')
            with self.assertRaises(ValueError):
                self.common.load_json(p)

    def test_formula_text_is_literal_and_numbers_remain_numbers(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'nested/safe.xlsx'
            self.common.write_xlsx(p, ['名稱', '金額'], [['=HYPERLINK("bad")', -10]])
            wb = load_workbook(p)
            self.assertEqual(wb.active['A2'].data_type, 's')
            self.assertEqual(wb.active['B2'].value, -10)
            self.assertTrue(wb.active.freeze_panes)

    def test_wrapped_excel_evidence_has_readable_row_height(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'readable.xlsx'
            self.common.write_xlsx(p, ['工作', '來源證據'], [['任務', '來源說明' * 22]])
            wb = load_workbook(p)
            self.assertGreater(wb.active.row_dimensions[2].height or 15, 96)

    def test_docx_and_text_roundtrip(self):
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'document.docx'
            self.common.write_docx(p, '模擬提案', [('目的', '保留已確認事實。'), ('待確認', ['負責人未提供'])])
            text = self.common.read_text(p)
            self.assertIn('模擬提案', text)
            self.assertIn('負責人未提供', text)

    def test_docx_extraction_preserves_paragraph_and_table_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'ordered.docx'
            doc = Document()
            doc.add_paragraph('已確認決議')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = '先內部試用'
            doc.add_paragraph('未定案提議')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = '全面上線'
            doc.save(p)
            text = self.common.read_text(p)
            self.assertLess(text.index('先內部試用'), text.index('未定案提議'))

    def test_unsupported_and_empty_source_rejected(self):
        with tempfile.TemporaryDirectory(dir=TEMP) as temp:
            p = Path(temp) / 'file.exe'
            p.write_bytes(b'fake')
            with self.assertRaises(ValueError):
                self.common.read_text(p)
            p = Path(temp) / 'empty.txt'
            p.write_text('', encoding='utf-8')
            with self.assertRaises(ValueError):
                self.common.read_text(p)


if __name__ == '__main__':
    unittest.main()
