"""實際產檔回歸；合成測資不代表真實公司資料驗收。"""
import json
import shutil
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]


class DailyWorkTests(unittest.TestCase):
    def test_research_report_describes_script_network_scope_without_claiming_url_origin(self):
        temp_root = ROOT / 'tests/.tmp'
        temp_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=temp_root) as temporary:
            folder = Path(temporary)
            (folder / 'source.md').write_text('服務包括一次檢查。', encoding='utf-8')
            (folder / 'input.json').write_text(json.dumps({'question': '包含什麼服務？', 'sources': [
                {'id': 's1', 'path': 'source.md', 'title': '實際讀取摘錄',
                 'url': 'https://example.org/service', 'accessed_at': '2026-11-03'}],
                'claims': [{'text': '包括一次檢查。', 'source_ids': ['s1'],
                            'excerpts': {'s1': '服務包括一次檢查。'}}]}, ensure_ascii=False), encoding='utf-8')
            result = subprocess.run([sys.executable, '-X', 'utf8',
                str(ROOT / '.agents/skills/research-assistant/scripts/run.py'),
                '--input', str(folder / 'input.json'), '--output', str(folder / 'result')],
                capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            report = (folder / 'result/research-report.md').read_text(encoding='utf-8')
            self.assertNotIn('使用者提供網址', report)
            self.assertNotIn('網址僅為使用者提供', report)
            self.assertIn('本程式沒有執行網路搜尋', report)
            self.assertFalse(json.loads((folder / 'result/evidence.json').read_text(encoding='utf-8'))['network_used'])

    def test_packaged_demo_inputs_run_from_only_the_installed_skill(self):
        temp_root = ROOT / 'tests/.tmp'
        temp_root.mkdir(parents=True, exist_ok=True)
        inputs = {'data-analyst': 'sales.csv', 'presentation-maker': 'deck.json', 'receipt-organizer': 'receipts.json'}
        with tempfile.TemporaryDirectory(dir=temp_root) as temporary:
            folder = Path(temporary)
            for skill in (ROOT / '.agents/skills').iterdir():
                if not (skill / 'SKILL.md').exists():
                    continue
                with self.subTest(skill=skill.name):
                    installed = folder / skill.name
                    shutil.copytree(skill, installed)
                    source = installed / 'references/demo' / inputs.get(skill.name, 'input.json')
                    self.assertTrue(source.is_file())
                    result = subprocess.run([sys.executable, '-X', 'utf8', str(installed / 'scripts/run.py'),
                        '--input', str(source), '--output', str(folder / (skill.name + '-out'))],
                        cwd=folder, capture_output=True, text=True, encoding='utf-8')
                    self.assertEqual(result.returncode, 0, result.stderr)
                    self.assertTrue(any((folder / (skill.name + '-out')).iterdir()))

    def test_receipt_explicit_non_simulation_does_not_request_simulation_choice(self):
        temp_root = ROOT / 'tests/.tmp'
        temp_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=temp_root) as temporary:
            folder = Path(temporary)
            source = folder / 'receipt.json'
            source.write_text(json.dumps({'source': '使用者提供收據', 'simulation': False,
                'receipts': [{'id': 'r1', 'date': None, 'vendor': '辦公用品店',
                    'items': '紙張', 'amount': '80.00', 'currency': 'TWD',
                    'category': '辦公用品', 'evidence': '來源第1列：紙張 TWD 80.00'}]},
                ensure_ascii=False), encoding='utf-8')
            result = subprocess.run([sys.executable, '-X', 'utf8', str(ROOT / '.agents/skills/receipt-organizer/scripts/run.py'),
                '--input', str(source), '--output', str(folder / 'result')], capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            report = (folder / 'result/receipts-report.md').read_text(encoding='utf-8')
            self.assertNotIn('是否模擬', report)
            self.assertIn('待人工確認', report)
            from openpyxl import load_workbook
            self.assertGreater(load_workbook(folder / 'result/receipts.xlsx').active.max_row, 1)
            self.assertFalse(json.loads((folder / 'result/receipts-summary.json').read_text(encoding='utf-8'))['simulation'])


if __name__ == '__main__':
    unittest.main()
