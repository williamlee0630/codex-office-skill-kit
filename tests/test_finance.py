"""財務教學模擬資料的金額與檔案行為驗收；預期值獨立列示。"""
import csv
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / '.agents/skills/finance-assistant'


class FinanceTests(unittest.TestCase):
    def setUp(self):
        temp_root = ROOT / 'tests/.tmp'
        temp_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=temp_root)
        self.work = Path(self.temp.name)
        self.out = self.work / 'out'

    def tearDown(self):
        self.temp.cleanup()

    def raw(self, name, rows, xlsx=False):
        path = self.work / name
        if xlsx:
            wb = Workbook(); ws = wb.active
            for row in rows: ws.append(row)
            wb.save(path)
        else:
            with path.open('w', encoding='utf-8-sig', newline='') as f:
                csv.writer(f).writerows(rows)
        return path

    def run_cli(self, data, ok=True, script=None, extra=()):
        if isinstance(data, dict):
            path = self.work / 'input.json'
            path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        else: path = Path(data)
        result = subprocess.run([sys.executable, '-X', 'utf8', str(script or SKILL / 'scripts/run.py'),
                                 '--input', str(path), '--output', str(self.out), *extra],
                                capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode == 0, ok, result.stdout + result.stderr)
        if ok:
            return json.loads((self.out / 'finance-result.json').read_text(encoding='utf-8'))
        self.assertIn('錯誤', result.stderr)
        return result

    def assert_common_sheets(self, filename):
        wb = load_workbook(self.out / filename, data_only=True)
        self.assertTrue({'原始資料', '計算依據', '待人工確認'} <= set(wb.sheetnames))
        self.assertTrue((self.out / '財務分析報告.docx').is_file())
        return wb

    def test_trial_balance_real_xlsx_numeric_and_unbalanced_no_fabrication(self):
        source = self.raw('tb.xlsx', [
            ['account', 'class', 'debit', 'credit', 'currency', 'period'],
            ['現金', 'asset', '1500.10', 0, 'TWD', '2026-10'],
            ['應付', 'liability', 0, '300.10', 'TWD', '2026-10'],
            ['資本', 'equity', 0, '1000', 'TWD', '2026-10'],
            ['營收', 'revenue', 0, '500', 'TWD', '2026-10'],
            ['費用', 'expense', '300', 0, 'TWD', '2026-10']], True)
        original = source.read_bytes()
        data = self.run_cli({'mode': 'statements', 'source': str(source)})
        group = data['groups'][0]
        self.assertEqual(group['profit'], '200.00')
        self.assertEqual(group['debit_total'], '1800.10')
        self.assertEqual(group['balance_difference'], '0.00')
        self.assertTrue(group['debit_credit_balanced'])
        self.assertFalse(group['cash_flow_available'])
        wb = self.assert_common_sheets('財務報表.xlsx')
        self.assertEqual(wb['損益表']['D2'].value, 500)
        self.assertEqual(wb['損益表']['D4'].value, 200)
        self.assertEqual(source.read_bytes(), original)
        shutil.rmtree(self.out)
        wb = load_workbook(source); wb.active['D4'] = 900; wb.save(source)
        data = self.run_cli({'mode': 'statements', 'source': str(source)})
        self.assertFalse(data['groups'][0]['debit_credit_balanced'])
        self.assertEqual(data['groups'][0]['balance_difference'], '100.00')
        self.assertIn('不平衡', '\n'.join(data['review']))

    def test_reconciliation_six_outcomes_reversals_one_to_one(self):
        header = ['date', 'reference', 'amount', 'currency', 'counterparty']
        left = self.raw('internal.csv', [header,
            ['2026-10-01', 'A', 100, 'TWD', '甲'],
            ['2026-10-02', 'B', 200, 'TWD', '乙'],
            ['2026-10-03', '', 300, 'TWD', '丙'],
            ['2026-10-04', 'ONLY-I', 400, 'TWD', '丁'],
            ['2026-10-05', 'DUP', 50, 'TWD', '戊'],
            ['2026-10-05', 'DUP', 50, 'TWD', '戊'],
            ['2026-10-06', 'REV', -25, 'TWD', '己']])
        right = self.raw('bank.csv', [header,
            ['2026-10-01', 'A', 100, 'TWD', '甲'],
            ['2026-10-02', 'B', 210, 'TWD', '乙'],
            ['2026-10-03', '', 300, 'TWD', '丙'],
            ['2026-10-04', 'ONLY-B', 600, 'TWD', '庚'],
            ['2026-10-06', 'REV', -25, 'TWD', '己']])
        data = self.run_cli({'mode': 'reconcile', 'source': str(left), 'other': str(right)})
        counts = data['counts']
        self.assertEqual(counts, {'已確認匹配': 2, '疑似匹配，待人工確認': 1,
                                 '金額不符': 1, '內部帳務缺漏': 1, '銀行或另一來源缺漏': 1,
                                 '疑似重複紀錄': 2})
        pairs = [r for r in data['results'] if r['internal_row'] and r['other_row']]
        self.assertEqual(len({r['other_row'] for r in pairs}), len(pairs))
        mismatch = next(r for r in pairs if r['status'] == '金額不符')
        self.assertEqual(mismatch['difference'], '-10.00')
        wb = self.assert_common_sheets('銀行對帳結果.xlsx')
        self.assertEqual(wb['分析結果']['G3'].value, -10)

    def test_amount_only_currency_period_date_and_missing_never_confirmed(self):
        header = ['date', 'reference', 'amount', 'currency', 'counterparty']
        a = self.raw('a.csv', [header, ['2026-10-01', '', 100, 'TWD', ''],
             ['2026-10-02', 'FX', 200, 'USD', '甲'], ['2026-10-02', 'MONTH', 500, 'TWD', '甲'],
             ['', 'NO-DATE', 20, 'TWD', '甲'], ['2026-10-02', 'BLANK', '', '', '']])
        b = self.raw('b.csv', [header, ['2026-10-01', '', 100, 'TWD', ''],
             ['2026-10-02', 'FX', 200, 'TWD', '甲'], ['2026-09-02', 'MONTH', 500, 'TWD', '甲'],
             ['2026-10-01', 'NO-DATE', 20, 'TWD', '甲']])
        data = self.run_cli({'mode': 'reconcile', 'source': str(a), 'other': str(b)})
        self.assertEqual(data['counts']['已確認匹配'], 0)
        self.assertTrue(any('空白' in s or '缺少' in s for s in data['review']))
        self.assertEqual(len(data['results']), 9)

    def test_variance_actual_base_zero_negative_decimal_currency_and_raw_mapping(self):
        source = self.raw('budget.txt', [
            ['項目', '預算', '實際', '幣別', '期間'],
            ['甲', '0', '0.105', 'TWD', '2026-10'],
            ['乙', '-100', '-80', 'TWD', '2026-10'],
            ['丙', '200', '250', 'USD', '2026-10'],
            ['丁', '', '12', 'USD', '2026-10']])
        data = self.run_cli({'mode': 'variance', 'source': str(source)})
        zero, negative, foreign = data['results'][:3]
        self.assertIsNone(zero['percentage'])
        self.assertEqual(zero['difference'], '0.11')
        self.assertEqual(negative['percentage'], '-20.00')
        self.assertEqual(foreign['percentage'], '25.00')
        wb = self.assert_common_sheets('財務差異分析.xlsx')
        self.assertEqual(wb['分析結果']['E2'].value, 0.11)
        self.assertIsNone(wb['分析結果']['F2'].value)
        self.assertEqual(wb['分析結果']['F4'].value, 25)
        self.assertEqual(len(data['totals']), 2)
        self.assertTrue(data['review'])

    def test_missing_account_and_period_preserved_without_complete_claim(self):
        source = self.raw('missing.csv', [['account', 'class', 'debit', 'credit', 'currency', 'period'],
                      ['', 'asset', '0.10', 0, 'TWD', '2026-10'],
                      ['存款', 'asset', '0.20', 0, 'TWD', ''],
                      ['資本', '', 0, '0.30', 'USD', '2026-10']])
        data = self.run_cli({'mode': 'statements', 'source': str(source)})
        self.assertFalse(data['complete'])
        self.assertTrue(data['review'])
        self.assertEqual(data['source_row_count'], 3)

    def test_two_source_variance_comparison_base_denominator_and_periods(self):
        headers = ['item', 'amount', 'currency', 'period']
        old = self.raw('old.csv', [headers, ['營收', 80, 'TWD', '2026-09']])
        new = self.raw('new.csv', [headers, ['營收', 100, 'TWD', '2026-10']])
        data = self.run_cli({'mode': 'variance', 'source': str(old), 'other': str(new),
                             'comparison': 'month'})
        self.assertEqual(data['results'][0]['percentage'], '25.00')
        self.assertEqual(data['results'][0]['base_period'], '2026-09')

    def test_rejects_invalid_nonfinite_formula_and_preserves_existing_outputs(self):
        source = self.raw('bad.csv', [['item', 'budget', 'actual', 'currency', 'period'],
                                    ['甲', 'NaN', '100', 'TWD', '2026-10']])
        self.run_cli({'mode': 'variance', 'source': str(source)}, ok=False)
        source = self.raw('ok.csv', [['item', 'budget', 'actual', 'currency', 'period'],
                                    ['=1+1', '100', '110', 'TWD', '2026-10']])
        self.run_cli({'mode': 'variance', 'source': str(source)})
        wb = self.assert_common_sheets('財務差異分析.xlsx')
        self.assertNotEqual(wb['原始資料']['D2'].data_type, 'f')
        original = (self.out / '財務差異分析.xlsx').read_bytes()
        self.run_cli({'mode': 'variance', 'source': str(source)}, ok=False)
        self.assertEqual((self.out / '財務差異分析.xlsx').read_bytes(), original)

    def test_excel_large_amount_preserves_exact_decimal_text(self):
        source = self.raw('precision.csv', [['item', 'budget', 'actual', 'currency', 'period'],
                       ['大額', '123456789012345.12', '123456789012345.23', 'TWD', '2026-10']])
        data = self.run_cli({'mode': 'variance', 'source': str(source)})
        self.assertEqual(data['results'][0]['difference'], '0.11')
        wb = self.assert_common_sheets('財務差異分析.xlsx')
        self.assertEqual(wb['分析結果']['C2'].value, '123456789012345.12')
        self.assertEqual(wb['分析結果']['D2'].value, '123456789012345.23')
        self.assertEqual(wb['分析結果']['E2'].value, 0.11)

    def test_text_source_physical_rows_survive_comments_blank_and_quoted_lines(self):
        source = self.work / 'physical.txt'
        source.write_text('# 教學模擬資料\n\nitem,budget,actual,currency,period\n"跨行\n項目",10,12,TWD,2026-10\n\n第二筆,20,22,TWD,2026-10\n', encoding='utf-8')
        data = self.run_cli({'mode': 'variance', 'source': str(source)})
        self.assertEqual([r['source_row'] for r in data['results']], [4, 7])
        wb = self.assert_common_sheets('財務差異分析.xlsx')
        self.assertIn([source.name, 4, 'item', '跨行\n項目'], [list(row)[:4] for row in wb['原始資料'].values])
        shutil.rmtree(self.out)
        source = self.work / 'physical.md'
        source.write_text('# 教學模擬資料\n\n| item | budget | actual | currency | period |\n|---|---|---|---|---|\n| 第一筆 | 10 | 12 | TWD | 2026-10 |\n\n| 第二筆 | 20 | 22 | TWD | 2026-10 |\n', encoding='utf-8')
        data = self.run_cli({'mode': 'variance', 'source': str(source)})
        self.assertEqual([r['source_row'] for r in data['results']], [5, 7])

    def test_csv_hash_item_is_data_and_three_decimal_currency_display(self):
        source = self.raw('hash.csv', [['item', 'budget', 'actual', 'currency', 'period'],
                       ['甲', '1.111', '1.234', 'KWD', '2026-10'],
                       ['#B', '200', '250', 'KWD', '2026-10']])
        data = self.run_cli({'mode': 'variance', 'source': str(source)})
        self.assertEqual(data['source_row_count'], 2)
        self.assertEqual(data['results'][1]['item'], '#B')
        wb = self.assert_common_sheets('財務差異分析.xlsx')
        self.assertEqual(wb['分析結果']['E2'].value, 0.123)
        self.assertEqual(wb['分析結果']['E2'].number_format, '#,##0.000;[Red]-#,##0.000')
        self.assertEqual(wb['分析結果']['K2'].number_format, '0')
        self.assertEqual(wb['分析結果']['F2'].number_format, '0.00"%"')

    def test_cash_flow_requires_complete_classification_and_closing_check(self):
        source = self.raw('cash.csv', [['account', 'class', 'debit', 'credit', 'currency', 'period',
                                      'cash_flow_class', 'cash_flow_amount'],
                  ['現金', 'asset', 100, 0, 'TWD', '2026-10', 'financing', 100],
                  ['資本', 'equity', 0, 100, 'TWD', '2026-10', 'noncash', 0]])
        config = {'mode': 'statements', 'source': str(source),
                  'cash_flow': {'opening_cash': '0', 'closing_cash': '100', 'currency': 'TWD', 'period': '2026-10'}}
        data = self.run_cli(config)
        self.assertTrue(data['groups'][0]['cash_flow_available'])
        wb = self.assert_common_sheets('財務報表.xlsx')
        self.assertEqual(wb['現金流量表']['D6'].value, 100)
        shutil.rmtree(self.out)
        config['cash_flow']['closing_cash'] = '101'
        data = self.run_cli(config)
        self.assertFalse(data['groups'][0]['cash_flow_available'])
        self.assertNotIn('現金流量表', load_workbook(self.out / '財務報表.xlsx').sheetnames)

    def test_requested_period_never_relabels_date_derived_source_period(self):
        a = self.raw('internal.csv', [['date', 'reference', 'amount', 'currency'], ['2026-09-10', 'A', 100, 'TWD']])
        b = self.raw('bank.csv', [['date', 'reference', 'amount', 'currency'], ['2026-09-10', 'A', 100, 'TWD']])
        result = self.run_cli({'mode': 'reconcile', 'source': str(a), 'other': str(b), 'period': '2026-10'})
        self.assertTrue(all(row['period'] == '2026-09' for row in result['results']))
        self.assertEqual(result['counts']['已確認匹配'], 0)
        self.assertIn('與指定會計期間不一致', '\n'.join(result['review']))

    def test_docx_missing_amounts_are_readable_unknowns(self):
        source = self.raw('missing.csv', [['item', 'budget', 'actual', 'currency', 'period'], ['空白測試', None, 10, 'TWD', '2026-10']])
        self.run_cli({'mode': 'variance', 'source': str(source)})
        from docx import Document
        report = '\n'.join(p.text for p in Document(self.out / '財務分析報告.docx').paragraphs)
        self.assertIn('待確認', report)
        self.assertNotIn('None', report)

    def test_three_simulated_fixtures_independent_expected_and_standalone(self):
        verify = subprocess.run([sys.executable, str(ROOT / 'examples/finance-assistant/verify_expected.py')],
                                capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(verify.returncode, 0, verify.stdout + verify.stderr)
        copied = self.work / 'standalone'
        shutil.copytree(SKILL, copied)
        for name in ['reconciliation', 'budget', 'trial-balance']:
            with self.subTest(name=name):
                self.out = self.work / name
                folder = ROOT / 'examples/finance-assistant' / name
                data = self.run_cli(folder / 'input.json', script=copied / 'scripts/run.py')
                expected = json.loads((folder / 'expected.json').read_text(encoding='utf-8'))
                if name == 'reconciliation': self.assertEqual(data['counts'], expected['counts'])
                elif name == 'budget':
                    self.assertEqual([r['difference'] for r in data['results']], expected['differences'])
                    self.assertEqual([r['percentage'] for r in data['results']], expected['percentages'])
                else:
                    for key, value in expected['group'].items():
                        self.assertEqual(data['groups'][0][key], value)
        self.out = self.work / 'default-demo'
        self.run_cli(copied / 'references/demo/input.json', script=copied / 'scripts/run.py')


if __name__ == '__main__':
    unittest.main()
