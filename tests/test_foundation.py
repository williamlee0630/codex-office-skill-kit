"""基礎技能行為測試；測試資料全部為模擬。"""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ("document-writer", "idea-generator")


class FoundationTests(unittest.TestCase):
    def execute(self, name, data=None):
        example = ROOT / "examples" / name / "input.json"
        if data is None:
            data = json.loads(example.read_text(encoding="utf-8"))
        temporary_root = ROOT / "tests/.tmp"
        temporary_root.mkdir(parents=True, exist_ok=True)
        temporary = tempfile.TemporaryDirectory(dir=temporary_root, prefix="foundation-")
        self.addCleanup(temporary.cleanup)
        folder = Path(temporary.name)
        input_path = folder / "input.json"
        input_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, "-X", "utf8", str(ROOT / ".agents/skills" / name / "scripts/run.py"), "--input", str(input_path), "--output", str(folder / "output")], capture_output=True, text=True, encoding="utf-8")
        return result, folder / "output"

    def test_normal_examples_produce_real_files(self):
        expected = {
            "document-writer": ["document-formal.md", "document-concise.md", "document-formal.docx", "document-concise.docx", "evidence.json"],
            "idea-generator": ["ideas.md", "comparison.csv", "selected-plan.md", "evidence.json"],
        }
        for name in SKILLS:
            with self.subTest(skill=name):
                result, output = self.execute(name)
                self.assertEqual(result.returncode, 0, result.stderr)
                for file_name in expected[name]:
                    self.assertGreater((output / file_name).stat().st_size, 10)
                    if file_name.endswith(".docx"):
                        from docx import Document
                        self.assertGreater(len(Document(output / file_name).paragraphs), 1)
                    if file_name.endswith(".xlsx"):
                        from openpyxl import load_workbook
                        self.assertGreater(load_workbook(output / file_name).active.max_row, 1)

    def test_incomplete_examples_keep_unknowns(self):
        for name in SKILLS:
            with self.subTest(skill=name):
                data = json.loads((ROOT / "examples" / name / "incomplete.json").read_text(encoding="utf-8"))
                result, output = self.execute(name, data)
                self.assertEqual(result.returncode, 0, result.stderr)
                content = "\n".join(path.read_text(encoding="utf-8-sig") for path in output.glob("*.md"))
                self.assertIn("待確認", content)

    def test_invalid_examples_fail_without_outputs(self):
        for name in SKILLS:
            with self.subTest(skill=name):
                data = json.loads((ROOT / "examples" / name / "invalid.json").read_text(encoding="utf-8"))
                result, output = self.execute(name, data)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("錯誤", result.stderr)
                self.assertFalse(output.exists())

    def test_wrong_format_rejected(self):
        for name in SKILLS:
            with self.subTest(skill=name):
                result, output = self.execute(name, ["不支援陣列根節點"])
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse(output.exists())

    def test_noncanonical_dates_are_rejected(self):
        for name in SKILLS:
            with self.subTest(skill=name):
                data = json.loads((ROOT / "examples" / name / "input.json").read_text(encoding="utf-8"))
                data["sources"][0]["date"] = "20261019"
                result, output = self.execute(name, data)
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse(output.exists())


    def test_unknown_source_rejected(self):
        for name in SKILLS:
            with self.subTest(skill=name):
                data = json.loads((ROOT / "examples" / name / "input.json").read_text(encoding="utf-8"))
                data["sources"] = [{"id": "other", "label": "模擬其他資料", "date": None}]
                result, output = self.execute(name, data)
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse(output.exists())





    def test_document_unknown_fact_ref_rejected(self):
        data = json.loads((ROOT / "examples/document-writer/input.json").read_text(encoding="utf-8"))
        data["versions"]["formal"][0]["fact_ids"] = ["missing"]
        result, output = self.execute("document-writer", data)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(output.exists())

    def test_ideas_default_requires_ten_distinct_ideas_and_subjective_label(self):
        result, output = self.execute("idea-generator")
        self.assertEqual(result.returncode, 0, result.stderr)
        content = (output / "ideas.md").read_text(encoding="utf-8")
        self.assertIn("主觀評估", content)
        data = json.loads((output / "evidence.json").read_text(encoding="utf-8"))
        self.assertEqual(len(data["ideas"]), 10)
        data["ideas"] = data["ideas"][:9]
        result, output = self.execute("idea-generator", data)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
