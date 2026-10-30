"""公開白名單、隱私、可攜式範例與實際 ZIP／Pages 連結的行為測試。"""
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
TEMP = ROOT / 'tests/.tmp'
TEMP.mkdir(parents=True, exist_ok=True)


def load_module(name):
    path = ROOT / 'tools' / (name + '.py')
    if not path.exists():
        raise AssertionError('尚未實作公開發行工具：' + name)
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def put(root, name, content='模擬公開資料\n'):
    path = root / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding='utf-8')
    return path


def fixture(root):
    for name in ['README.md', 'ACCEPTANCE.md', 'AGENTS.md', 'LICENSE', 'requirements.txt',
                 '.gitignore', 'docs/INSTALL.md', 'docs/ACCEPTANCE.md', 'qa/summary.md']:
        put(root, name)
    put(root, 'index.html', '<meta http-equiv="refresh" content="0; url=handbook/index.html"><a href="handbook/index.html">講義</a>')
    put(root, 'handbook/index.html', '<a href="../dist/codex-office-skill-kit-1.1.0.zip">完整學員工具包</a>'
        '<a href="../dist/office-skills-1.1.0.zip">技能</a><a href="../dist/practice-data-1.1.0.zip">資料</a>'
        '<a href="../examples/demo/source.md">來源</a><a href="../.agents/skills/demo-0/SKILL.md">技能原文</a>'
        '<script src="assets/app.js"></script>')
    put(root, 'handbook/assets/app.js', '"use strict";')
    put(root, 'handbook/build.py', 'raise RuntimeError("不應出現在 Pages")')
    put(root, 'examples/demo/source.md')
    for i in range(10):
        put(root, f'.agents/skills/demo-{i}/SKILL.md')
    put(root, 'tools/office_common.py', 'pass\n')
    put(root, 'tests/test_common.py', 'import unittest\n')
    put(root, 'tests/test_publication.py', 'raise RuntimeError("精簡包不應載入")')
    put(root, 'tests/test_handbook_learner.py', 'raise RuntimeError("精簡包不應載入")')
    return root


class PublicationTests(unittest.TestCase):
    def test_skills_package_has_no_links_to_missing_full_handbook_or_qa(self):
        release = load_module('build_release')
        checker = load_module('check_release')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = fixture(Path(tmp))
            put(root, 'ACCEPTANCE.md', '[驗收](qa/summary.md)')
            put(root, 'docs/INSTALL.md', '[官方查證](OFFICIAL_SOURCES.md)')
            put(root, 'docs/OFFICIAL_SOURCES.md', '官方資料')
            put(root, 'docs/PROGRESS.md', '[講義](../handbook/index.html)')
            put(root, 'tests/manual/trigger-cases.md', '[驗收](../../ACCEPTANCE.md)')
            release.build(root)
            with ZipFile(root / 'dist/office-skills-1.1.0.zip') as bundle:
                files = {name: bundle.read(name) for name in bundle.namelist()}
            self.assertGreater(checker._check_documents(files), 0)
            self.assertIn('docs/INSTALL.md', files)
            self.assertIn('docs/OFFICIAL_SOURCES.md', files)

    def test_whitelist_excludes_internal_qa_credentials_and_temp(self):
        publication = load_module('publication')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = fixture(Path(tmp))
            for name in ['docs/qa/events.json', '.env', '.codex/config.toml', '.aws/credentials',
                         'examples/demo/.env.production', 'examples/demo/local-settings.json',
                         'examples/demo/credentials.yaml', 'qa/raw-events.json',
                         'tests/.tmp/private.json', '.work/cache.md', '.venv/readme.md',
                         'qa/raw-events.log', 'dist/old.zip', 'examples/demo/~$office.docx']:
                put(root, name, '私人資料')
            selected = {p.relative_to(root).as_posix() for p in publication.public_files(root)}
            self.assertIn('qa/summary.md', selected)
            self.assertIn('.agents/skills/demo-0/SKILL.md', selected)
            self.assertIn('handbook/build.py', selected)  # 原始碼可公開；Pages 另有精簡白名單。
            self.assertFalse(any('docs/qa/' in p or '/.tmp/' in p or '.env' in p or
                                 'local-settings' in p or 'raw-events' in p or 'old.zip' in p or '~$' in p for p in selected))
            self.assertNotIn('examples/demo/credentials.yaml', selected)

    def test_scan_rejects_home_email_token_and_nested_zip(self):
        publication = load_module('publication')
        home = 'C:' + chr(92) + 'Users' + chr(92) + 'private-person' + chr(92) + 'secret.txt'
        private_mail = 'someone' + '@' + 'private-company.test'
        token = 'gh' + 'p_' + 'a' * 36
        for value in [home, private_mail, token, '/ho' + 'me/person/private.txt']:
            with self.subTest(kind=value[:4]):
                self.assertTrue(publication.scan(value.encode(), label='模擬測試'))
        self.assertEqual(publication.scan(b'a@example.invalid b@example.com c@example.org 12+demo@users.noreply.github.com'), [])
        inner = io.BytesIO()
        with ZipFile(inner, 'w') as archive:
            archive.writestr('xl/sharedStrings.xml', home)
        outer = io.BytesIO()
        with ZipFile(outer, 'w') as archive:
            archive.writestr('nested/sample.xlsx', inner.getvalue())
        issues = publication.scan(outer.getvalue(), label='outer.zip')
        self.assertTrue(any('sharedStrings.xml' in issue for issue in issues))
        self.assertFalse(any('private-person' in issue for issue in issues), '診斷不得再次洩露私密值')

    def test_scan_checks_utf16_and_keeps_private_archive_names_out_of_diagnostics(self):
        publication = load_module('publication')
        home = 'C:' + chr(92) + 'Users' + chr(92) + 'sensitive-person' + chr(92) + 'document.xml'
        self.assertTrue(publication.scan(home.encode('utf-16'), label='模擬 UTF16'))
        archive_bytes = io.BytesIO()
        with ZipFile(archive_bytes, 'w') as archive:
            archive.writestr(home, '模擬資料')
        issues = publication.scan(archive_bytes.getvalue(), label='模擬 ZIP')
        self.assertTrue(issues)
        self.assertFalse(any('sensitive-person' in issue for issue in issues))

    def test_zip_scan_checks_expanded_content_instead_of_compressed_noise(self):
        publication = load_module('publication')
        data = (ROOT / 'tools/office_common.py').read_bytes()
        self.assertEqual(publication.scan(data), [])
        archive_bytes = io.BytesIO()
        with ZipFile(archive_bytes, 'w', ZIP_DEFLATED, compresslevel=9) as archive:
            archive.writestr('office_common.py', data)
        self.assertEqual(publication.scan(archive_bytes.getvalue()), [])

    def test_zip_rejects_windows_parent_traversal(self):
        publication = load_module('publication')
        archive_bytes = io.BytesIO()
        with ZipFile(archive_bytes, 'w') as archive:
            archive.writestr('safe/' + '..' + chr(92) + 'outside.txt', '模擬資料')
        self.assertTrue(publication.scan(archive_bytes.getvalue()))

    def test_sanitize_preserves_source_numbers_xlsx_and_original_backup(self):
        publication = load_module('publication')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = Path(tmp).resolve()
            source = put(root, 'examples/demo/source.md')
            record = put(root, 'examples/demo/output/evidence.json', json.dumps({'source': str(source), 'amount': 120.5}))
            outside = put(root, 'personal.json', json.dumps({'source': str(source)}))
            book = root / 'examples/demo/output/result.xlsx'
            workbook = Workbook()
            workbook.active.append(['來源', '金額'])
            workbook.active.append([str(source), 120.5])
            workbook.save(book)
            old_record, old_book = record.read_bytes(), book.read_bytes()
            changed = publication.sanitize_examples(root)
            self.assertEqual(set(changed), {'examples/demo/output/evidence.json', 'examples/demo/output/result.xlsx'})
            self.assertEqual(json.loads(record.read_text(encoding='utf-8')), {'source': 'examples/demo/source.md', 'amount': 120.5})
            opened = load_workbook(book)
            self.assertEqual(opened.active['A2'].value, 'examples/demo/source.md')
            self.assertEqual(opened.active['B2'].value, 120.5)
            self.assertEqual(publication.scan(book), [])
            self.assertEqual((root / '.work/publication-backup/examples/demo/output/evidence.json').read_bytes(), old_record)
            self.assertEqual((root / '.work/publication-backup/examples/demo/output/result.xlsx').read_bytes(), old_book)
            self.assertEqual(json.loads(outside.read_text(encoding='utf-8'))['source'], str(source))
            self.assertEqual(publication.sanitize_examples(root), [], '重跑不改寫也不覆寫備份')

    def test_release_and_pages_resolve_links_under_repo_subpath(self):
        release = load_module('build_release')
        checker = load_module('check_release')
        site = load_module('build_site')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = fixture(Path(tmp))
            manifest = release.build(root)
            self.assertEqual(len(manifest['packages']), 3)
            checked = checker.check(root, write_report=False)
            self.assertEqual(len(checked), 3)
            destination = site.build(root)
            self.assertTrue((destination / '.nojekyll').exists())
            self.assertFalse((destination / 'handbook/build.py').exists())
            self.assertFalse((destination / 'tests').exists())
            # upload-pages-artifact 會排除點開頭目錄；網站另用公開別名，離線包維持原結構。
            self.assertTrue((destination / 'skills/demo-0/SKILL.md').exists())
            self.assertFalse((destination / '.agents').exists())
            public_html = (destination / 'handbook/index.html').read_text(encoding='utf-8')
            self.assertIn('href="../skills/demo-0/SKILL.md"', public_html)
            self.assertGreater(checker.check_links(destination), 5)
            archive = root / 'dist/codex-office-skill-kit-1.1.0.zip'
            with ZipFile(archive) as bundle:
                html = bundle.read('codex-office-skill-kit/handbook/index.html').decode()
                self.assertNotIn('href="../dist/codex-office-skill-kit-1.1.0.zip"', html)
                self.assertIn('href="../README.md"', html)
                self.assertIn('codex-office-skill-kit/ACCEPTANCE.md', bundle.namelist())
                self.assertIn('codex-office-skill-kit/qa/summary.md', bundle.namelist())
                self.assertIn('href="../.agents/skills/demo-0/SKILL.md"', html)
            with ZipFile(root / 'dist/office-skills-1.1.0.zip') as bundle:
                self.assertNotIn('tests/test_publication.py', bundle.namelist())
                self.assertNotIn('tests/test_handbook_learner.py', bundle.namelist())

    def test_selected_private_content_missing_download_and_tamper_stop_publication(self):
        release = load_module('build_release')
        checker = load_module('check_release')
        site = load_module('build_site')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = fixture(Path(tmp))
            private = put(root, 'docs/qa/private.json', 'gh' + 'p_' + 'x' * 36)
            release.build(root)  # 未選取的歷史 QA 不使公開建置失敗。
            selected = put(root, 'qa/summary.md', 'owner' + '@' + 'private.test')
            with self.assertRaises(ValueError):
                release.build(root)
            selected.write_text('模擬驗收摘要', encoding='utf-8')
            release.build(root)
            (root / 'dist/practice-data-1.1.0.zip').unlink()
            with self.assertRaises((ValueError, FileNotFoundError)):
                site.build(root)
            release.build(root)
            with (root / 'dist/practice-data-1.1.0.zip').open('ab') as handle:
                handle.write(b'tampered')
            with self.assertRaises(ValueError):
                checker.check(root, write_report=False)
            release.build(root)
            (root / 'dist/codex-office-skill-kit-1.1.0.zip').unlink()
            manifest_path = root / 'dist/manifest.json'
            manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
            manifest['packages'] = manifest['packages'][:2]
            manifest_path.write_text(json.dumps(manifest), encoding='utf-8')
            with self.assertRaises(ValueError):
                checker.check(root, write_report=False)
            (root / 'ACCEPTANCE.md').unlink()
            with self.assertRaisesRegex(ValueError, 'ACCEPTANCE'):
                release.build(root)

    def test_local_link_checker_rejects_missing_assets_and_site_absolute_urls(self):
        checker = load_module('check_release')
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            root = Path(tmp)
            page = put(root, 'index.html', '<script src="missing.js"></script>')
            with self.assertRaises(ValueError):
                checker.check_links(root)
            put(root, 'exists.html', '<p>模擬資料</p>')
            page.write_text('<a href="/exists.html">連結</a>', encoding='utf-8')
            with self.assertRaises(ValueError):
                checker.check_links(root)


if __name__ == '__main__':
    unittest.main()
